#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
HTML slides -> editable PPTX converter.

Dependencies (install via pip):
  pip install beautifulsoup4 python-pptx pillow playwright
  playwright install chromium  # 初回のみ

Usage:
  python html_ppt.py input.html output.pptx --mode editable

The script reads a single HTML document, interprets elements with class="slide"
as individual slides (falling back to <body> when not provided), and recreates
text, lists, tables, images, and simple shapes as native PowerPoint objects.
The focus is on editable output rather than pixel-perfect fidelity.
"""

from __future__ import annotations

import argparse
import base64
import io
import math
import re
import sys
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple
from urllib.parse import unquote, urlparse
from urllib.request import url2pathname

from bs4 import BeautifulSoup, NavigableString, Tag
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sync_playwright = None
_PLAYWRIGHT_IMPORT_ATTEMPTED = False
Image = None
ImageColor = None
_PIL_IMPORT_ATTEMPTED = False

# Slide canvas size used internally (px). Actual PPTX is scaled proportionally.
SLIDE_REF_WIDTH = 1920
SLIDE_REF_HEIGHT = 1080
DEFAULT_PADDING_X = 140
DEFAULT_PADDING_Y = 120
FLOW_GAP = 40
DEFAULT_BLOCK_HEIGHT = 140
DEFAULT_LIST_INDENT = 60

# Default font sizes (px) per tag for fallback cases.
DEFAULT_FONT_SIZE = 28
TAG_FONT_SIZE = {
    "h1": 64,
    "h2": 48,
    "h3": 36,
    "h4": 30,
    "p": 28,
    "li": 26,
}

CSS_COMMENT_RE = re.compile(r"/\*.*?\*/", re.S)
SIMPLE_SELECTOR_TOKEN_RE = re.compile(r"([#.]?[\w-]+|\*)")
WHITESPACE_RE = re.compile(r"\s+")
GRID_REPEAT_RE = re.compile(r"repeat\((\d+),\s*([^)]+)\)")
CSS_COLOR_KEYWORDS = {
    "white": (255, 255, 255),
    "black": (0, 0, 0),
    "red": (255, 0, 0),
    "green": (0, 128, 0),
    "lime": (0, 255, 0),
    "blue": (0, 0, 255),
    "yellow": (255, 255, 0),
    "orange": (255, 165, 0),
    "purple": (128, 0, 128),
    "gray": (128, 128, 128),
    "grey": (128, 128, 128),
    "silver": (192, 192, 192),
    "navy": (0, 0, 128),
    "teal": (0, 128, 128),
    "cyan": (0, 255, 255),
    "aqua": (0, 255, 255),
    "magenta": (255, 0, 255),
    "fuchsia": (255, 0, 255),
    "maroon": (128, 0, 0),
    "olive": (128, 128, 0),
}

GENERIC_FONT_MAPPING = {
    "sans-serif": "Aptos",
    "serif": "Times New Roman",
    "monospace": "Consolas",
    "system-ui": "Aptos",
    "ui-sans-serif": "Aptos",
    "ui-serif": "Times New Roman",
    "ui-monospace": "Consolas",
}

PPT_PAGE_SIZES_INCHES = {
    "16:9": (13.333, 7.5),
    "a4": (11.693, 8.268),  # A4 landscape
}

ImageMap = Dict[str, str]


def ensure_playwright():
    global sync_playwright, _PLAYWRIGHT_IMPORT_ATTEMPTED
    if sync_playwright is None and not _PLAYWRIGHT_IMPORT_ATTEMPTED:
        _PLAYWRIGHT_IMPORT_ATTEMPTED = True
        try:
            from playwright.sync_api import sync_playwright as imported_sync_playwright
        except ImportError:
            return None
        sync_playwright = imported_sync_playwright
    return sync_playwright


def ensure_pillow():
    global Image, ImageColor, _PIL_IMPORT_ATTEMPTED
    if Image is None and ImageColor is None and not _PIL_IMPORT_ATTEMPTED:
        _PIL_IMPORT_ATTEMPTED = True
        try:
            from PIL import Image as imported_image, ImageColor as imported_image_color
        except ImportError:
            return None, None
        Image = imported_image
        ImageColor = imported_image_color
    return Image, ImageColor


def normalize_page_size(value: str) -> str:
    normalized = (value or "").strip().lower()
    if normalized in {"16:9", "16x9", "widescreen"}:
        return "16:9"
    if normalized in {"a4", "a4-landscape", "a4_l"}:
        return "a4"
    raise argparse.ArgumentTypeError("ページサイズは '16:9' または 'A4' を指定してください。")


def parse_image_map_entry(value: str) -> Tuple[str, str]:
    key, sep, mapped_path = value.partition("=")
    if not sep or not key.strip() or not mapped_path.strip():
        raise argparse.ArgumentTypeError("画像マッピングは 'IMAGE_URL_1=/path/to/file.png' 形式で指定してください。")
    return key.strip(), mapped_path.strip()


def encode_png_data_url(image_bytes: bytes) -> str:
    return "data:image/png;base64," + base64.b64encode(image_bytes).decode("ascii")


def has_large_column_panels(slide_model: "SlideModel") -> bool:
    canvas_width = slide_model.canvas_width or SLIDE_REF_WIDTH
    canvas_height = slide_model.canvas_height or SLIDE_REF_HEIGHT
    panels = [
        block
        for block in slide_model.blocks
        if block.kind == "shape"
        and (block.layout.width or 0) >= canvas_width * 0.28
        and (block.layout.height or 0) >= canvas_height * 0.6
        and (block.layout.top or 0) <= canvas_height * 0.25
    ]
    return len(panels) >= 2


def has_wide_bottom_callout(slide_model: "SlideModel") -> bool:
    canvas_width = slide_model.canvas_width or SLIDE_REF_WIDTH
    canvas_height = slide_model.canvas_height or SLIDE_REF_HEIGHT
    for block in slide_model.blocks:
        if block.kind != "shape":
            continue
        width = block.layout.width or 0
        height = block.layout.height or 0
        top = block.layout.top or 0
        if width >= canvas_width * 0.85 and canvas_height * 0.45 <= top <= canvas_height * 0.75 and 36 <= height <= canvas_height * 0.2:
            return True
    return False


def detect_auto_rasterize_slides(slides: List["SlideModel"]) -> Dict[int, str]:
    targets: Dict[int, str] = {}
    for index, slide_model in enumerate(slides):
        block_kinds = [block.kind for block in slide_model.blocks]
        table_count = block_kinds.count("table")
        if not table_count:
            continue
        if "bar-chart" in block_kinds:
            targets[index] = "table+chart"
            continue
        if has_large_column_panels(slide_model) and (block_kinds.count("shape") >= 4 or block_kinds.count("list") >= 1):
            targets[index] = "dense-two-column"
            continue
        if has_wide_bottom_callout(slide_model):
            targets[index] = "wide-callout"
    return targets


def resolve_rasterize_slide_targets(spec: str, slides: List["SlideModel"]) -> Dict[int, str]:
    normalized = (spec or "").strip().lower()
    if normalized in {"", "none", "off", "false"}:
        return {}
    targets: Dict[int, str] = {}
    if normalized == "auto":
        return detect_auto_rasterize_slides(slides)
    auto_targets: Optional[Dict[int, str]] = None
    for token in (part.strip() for part in normalized.split(",")):
        if not token:
            continue
        if token == "auto":
            if auto_targets is None:
                auto_targets = detect_auto_rasterize_slides(slides)
            targets.update(auto_targets)
            continue
        if "-" in token:
            start_str, end_str = token.split("-", 1)
            if not start_str.isdigit() or not end_str.isdigit():
                raise ValueError(f"無効なスライド範囲です: {token}")
            start = int(start_str)
            end = int(end_str)
            if start <= 0 or end <= 0:
                raise ValueError(f"スライド番号は 1 以上で指定してください: {token}")
            if end < start:
                start, end = end, start
            for slide_number in range(start, end + 1):
                if slide_number <= len(slides):
                    targets[slide_number - 1] = "manual"
            continue
        if not token.isdigit():
            raise ValueError(f"無効なスライド指定です: {token}")
        slide_number = int(token)
        if slide_number <= 0:
            raise ValueError(f"スライド番号は 1 以上で指定してください: {token}")
        if slide_number <= len(slides):
            targets[slide_number - 1] = "manual"
    return targets

BROWSER_COLLECT_JS = """
({ selector }) => {
  const TEXT_ACCEPT = new Set(["h1","h2","h3","h4","h5","h6","p","blockquote","div","section","article"]);
  const BLOCK_TAGS = new Set(["p","div","section","article","ul","ol","table","li","h1","h2","h3","h4","h5","h6"]);

  const slides = Array.from(document.querySelectorAll(selector)).filter(el => el.offsetWidth > 0 && el.offsetHeight > 0);
  const targets = slides.length ? slides : [document.body];

  const parseZ = (value) => {
    if (!value || value === "auto") return 0;
    const n = Number(value);
    return Number.isFinite(n) ? n : 0;
  };

  const hasNestedBlocks = (node) => {
    for (const child of Array.from(node.children || [])) {
      const tag = child.tagName ? child.tagName.toLowerCase() : "";
      if (BLOCK_TAGS.has(tag) && child.innerText && child.innerText.trim().length) {
        return true;
      }
    }
    return false;
  };

  const normalizeSpaces = (text) => {
    if (!text) return "";
    return text.replace(/\\u00a0/g, " ").replace(/\\s+/g, " ");
  };

  const applyTextTransform = (text, transform) => {
    if (!text) return "";
    const normalized = (transform || "").toLowerCase();
    if (!normalized || normalized === "none") return text;
    if (normalized === "uppercase") return text.toUpperCase();
    if (normalized === "lowercase") return text.toLowerCase();
    if (normalized === "capitalize") {
      return text.replace(/(^|[\\s\\u3000])(\\S)/g, (match, prefix, char) => `${prefix}${char.toUpperCase()}`);
    }
    return text;
  };

  const mergeRuns = (runs) => {
    const merged = [];
    for (const run of runs) {
      if (!run.text) continue;
      const prev = merged[merged.length - 1];
      if (
        prev &&
        prev.fontSizePx === run.fontSizePx &&
        prev.fontWeight === run.fontWeight &&
        prev.fontStyle === run.fontStyle &&
        prev.color === run.color &&
        prev.fontFamily === run.fontFamily
      ) {
        prev.text += run.text;
      } else {
        merged.push({ ...run });
      }
    }
    return merged;
  };

  const collectInlineRuns = (root, skipNestedBlockNodes = false) => {
    const baseStyle = getComputedStyle(root);
    const base = {
      fontSizePx: parseFloat(baseStyle.fontSize) || undefined,
      fontWeight: baseStyle.fontWeight,
      fontStyle: baseStyle.fontStyle,
      color: baseStyle.color,
      fontFamily: baseStyle.fontFamily,
      textTransform: baseStyle.textTransform,
    };
    const rawRuns = [];
    const traverse = (node, style) => {
      if (node.nodeType === Node.TEXT_NODE) {
        const normalized = applyTextTransform(normalizeSpaces(node.nodeValue || ""), style.textTransform);
        if (!normalized) return;
        rawRuns.push({
          text: normalized,
          fontSizePx: style.fontSizePx,
          fontWeight: style.fontWeight,
          fontStyle: style.fontStyle,
          color: style.color,
          fontFamily: style.fontFamily,
          textTransform: style.textTransform,
        });
        return;
      }
      if (node.nodeType !== Node.ELEMENT_NODE) return;
      const tag = node.tagName ? node.tagName.toLowerCase() : "";
      if (tag === "script" || tag === "style" || tag === "noscript") return;
      if (tag === "br") {
        rawRuns.push({
          text: "\\n",
          fontSizePx: style.fontSizePx,
          fontWeight: style.fontWeight,
          fontStyle: style.fontStyle,
          color: style.color,
          fontFamily: style.fontFamily,
          textTransform: style.textTransform,
        });
        return;
      }
      if (skipNestedBlockNodes && BLOCK_TAGS.has(tag) && tag !== "span") {
        return;
      }
      const cs = getComputedStyle(node);
      const next = { ...style };
      if (cs.fontSize) next.fontSizePx = parseFloat(cs.fontSize) || next.fontSizePx;
      if (cs.fontWeight) next.fontWeight = cs.fontWeight || next.fontWeight;
      if (cs.fontStyle) next.fontStyle = cs.fontStyle || next.fontStyle;
      if (cs.color) next.color = cs.color || next.color;
      if (cs.fontFamily) next.fontFamily = cs.fontFamily || next.fontFamily;
      if (cs.textTransform) next.textTransform = cs.textTransform || next.textTransform;
      if (tag === "strong" || tag === "b") next.fontWeight = "700";
      if (tag === "em" || tag === "i") next.fontStyle = "italic";
      for (const child of Array.from(node.childNodes || [])) {
        traverse(child, next);
      }
    };
    for (const child of Array.from(root.childNodes || [])) {
      traverse(child, base);
    }
    return mergeRuns(rawRuns).filter(run => run.text && run.text.length);
  };

  const isExcluded = (node, excludedElements) => {
    if (!excludedElements || !excludedElements.size || !node) return false;
    let current = node.nodeType === Node.ELEMENT_NODE ? node : node.parentElement;
    while (current && current.nodeType === Node.ELEMENT_NODE) {
      if (excludedElements.has(current)) return true;
      current = current.parentElement;
    }
    return false;
  };

  const parseNumericText = (text) => {
    if (!text) return null;
    const match = text.replace(/,/g, "").match(/-?\d+(?:\.\d+)?/);
    if (!match) return null;
    const value = parseFloat(match[0]);
    return Number.isFinite(value) ? value : null;
  };

  const parseCssColor = (value) => {
    if (!value) return null;
    const rgba = value.match(/rgba?\(([^)]+)\)/i);
    if (rgba) {
      const parts = rgba[1].split(",").map((part) => parseFloat(part.trim()));
      if (parts.length >= 3) {
        return [
          Math.max(0, Math.min(255, Math.round(parts[0]))),
          Math.max(0, Math.min(255, Math.round(parts[1]))),
          Math.max(0, Math.min(255, Math.round(parts[2]))),
          parts.length >= 4 && Number.isFinite(parts[3]) ? Math.max(0, Math.min(1, parts[3])) : 1,
        ];
      }
    }
    const hex = value.match(/#([0-9a-f]{3,8})/i);
    if (!hex) return null;
    const raw = hex[1];
    if (raw.length === 3) {
      return [
        parseInt(raw[0] + raw[0], 16),
        parseInt(raw[1] + raw[1], 16),
        parseInt(raw[2] + raw[2], 16),
        1,
      ];
    }
    if (raw.length >= 6) {
      return [
        parseInt(raw.slice(0, 2), 16),
        parseInt(raw.slice(2, 4), 16),
        parseInt(raw.slice(4, 6), 16),
        raw.length >= 8 ? parseInt(raw.slice(6, 8), 16) / 255 : 1,
      ];
    }
    return null;
  };

  const approximateLinearGradientColor = (bgImage) => {
    if (!bgImage || !bgImage.includes("linear-gradient")) return null;
    const matches = bgImage.match(/rgba?\([^)]*\)|#[0-9a-fA-F]{3,8}/g);
    if (!matches || !matches.length) return null;
    const parsed = matches.map(parseCssColor).filter(Boolean);
    if (!parsed.length) return null;
    const first = parsed[0];
    const last = parsed[parsed.length - 1];
    const blend = [
      Math.round((first[0] + last[0]) / 2),
      Math.round((first[1] + last[1]) / 2),
      Math.round((first[2] + last[2]) / 2),
      (first[3] + last[3]) / 2,
    ];
    return `rgba(${blend[0]}, ${blend[1]}, ${blend[2]}, ${blend[3].toFixed(3)})`;
  };

  const collectRichTextRuns = (root) => {
    const baseStyle = getComputedStyle(root);
    const base = {
      fontSizePx: parseFloat(baseStyle.fontSize) || undefined,
      fontWeight: baseStyle.fontWeight,
      fontStyle: baseStyle.fontStyle,
      color: baseStyle.color,
      fontFamily: baseStyle.fontFamily,
      textTransform: baseStyle.textTransform,
    };
    const rawRuns = [];
    const pushRun = (text, style) => {
      if (!text) return;
      rawRuns.push({
        text,
        fontSizePx: style.fontSizePx,
        fontWeight: style.fontWeight,
        fontStyle: style.fontStyle,
        color: style.color,
        fontFamily: style.fontFamily,
        textTransform: style.textTransform,
      });
    };
    const trimTrailingWhitespace = () => {
      while (rawRuns.length) {
        const last = rawRuns[rawRuns.length - 1];
        if (!last) break;
        if (last.text === "\\n") return;
        const trimmed = last.text.replace(/[ \t]+$/g, "");
        if (trimmed === "") {
          rawRuns.pop();
          continue;
        }
        last.text = trimmed;
        return;
      }
    };
    const ensureNewline = (style) => {
      trimTrailingWhitespace();
      const last = rawRuns[rawRuns.length - 1];
      if (!last || last.text.endsWith("\\n")) return;
      pushRun("\\n", style);
    };
    const traverse = (node, style) => {
      if (node.nodeType === Node.TEXT_NODE) {
        const rawText = node.nodeValue || "";
        if (!/\\S/.test(rawText)) return;
        const normalized = applyTextTransform(normalizeSpaces(rawText), style.textTransform);
        if (normalized) pushRun(normalized, style);
        return;
      }
      if (node.nodeType !== Node.ELEMENT_NODE) return;
      const tag = node.tagName ? node.tagName.toLowerCase() : "";
      if (tag === "script" || tag === "style" || tag === "noscript") return;
      if (tag === "br") {
        pushRun("\\n", style);
        return;
      }
      const cs = getComputedStyle(node);
      const next = { ...style };
      if (cs.fontSize) next.fontSizePx = parseFloat(cs.fontSize) || next.fontSizePx;
      if (cs.fontWeight) next.fontWeight = cs.fontWeight || next.fontWeight;
      if (cs.fontStyle) next.fontStyle = cs.fontStyle || next.fontStyle;
      if (cs.color) next.color = cs.color || next.color;
      if (cs.fontFamily) next.fontFamily = cs.fontFamily || next.fontFamily;
      if (cs.textTransform) next.textTransform = cs.textTransform || next.textTransform;
      if (tag === "strong" || tag === "b") next.fontWeight = "700";
      if (tag === "em" || tag === "i") next.fontStyle = "italic";
      if (tag === "li") {
        ensureNewline(next);
        pushRun("• ", next);
      } else if (["p", "div"].includes(tag) && rawRuns.length) {
        ensureNewline(next);
      }
      for (const child of Array.from(node.childNodes || [])) {
        traverse(child, next);
      }
      if (["li", "p", "div"].includes(tag)) {
        ensureNewline(next);
      }
    };
    for (const child of Array.from(root.childNodes || [])) {
      traverse(child, base);
    }
    const merged = mergeRuns(rawRuns)
      .map((run) => ({
        ...run,
        text: run.text
          .replace(/[ \t]*\\n[ \t]*/g, "\\n")
          .replace(/\\n{2,}/g, "\\n"),
      }))
      .filter((run) => run.text && run.text.replace(/\\n/g, "").trim().length);
    if (merged.length) {
      merged[0].text = merged[0].text.replace(/^\\n+/, "");
      merged[merged.length - 1].text = merged[merged.length - 1].text.replace(/\\n+$/, "");
    }
    return merged.filter(run => run.text && run.text.length);
  };

  const collectTextBlocks = (slide, slideRect, counter, excludedElements = null) => {
    const blocks = [];
    const walker = document.createTreeWalker(slide, NodeFilter.SHOW_ELEMENT, {
      acceptNode(node) {
        if (isExcluded(node, excludedElements)) return NodeFilter.FILTER_REJECT;
        const tag = node.tagName ? node.tagName.toLowerCase() : "";
        if (!TEXT_ACCEPT.has(tag)) return NodeFilter.FILTER_SKIP;
        if (tag === "div" || tag === "section" || tag === "article" || tag === "span") {
          if (!node.innerText || !node.innerText.trim()) return NodeFilter.FILTER_SKIP;
          if (hasNestedBlocks(node)) return NodeFilter.FILTER_SKIP;
        }
        if (node.closest("ul,ol,table")) return NodeFilter.FILTER_SKIP;
        return NodeFilter.FILTER_ACCEPT;
      }
    });
    while (walker.nextNode()) {
      const el = walker.currentNode;
      const rect = el.getBoundingClientRect();
      if (rect.width < 4 || rect.height < 4) continue;
      const cs = getComputedStyle(el);
      const runs = collectInlineRuns(el);
      if (!runs.length) continue;
      const combinedText = runs.map(run => run.text).join("");
      const zIndex = parseZ(cs.zIndex);
      const block = {
        kind: "text",
        rect: {
          x: rect.x - slideRect.x,
          y: rect.y - slideRect.y,
          w: rect.width,
          h: rect.height
        },
        text: combinedText,
        runs,
        styles: {
          fontSizePx: parseFloat(cs.fontSize) || undefined,
          lineHeightPx: parseFloat(cs.lineHeight) || undefined,
          fontWeight: cs.fontWeight,
          fontStyle: cs.fontStyle,
          color: cs.color,
          fontFamily: cs.fontFamily,
          textAlign: cs.textAlign,
          textTransform: cs.textTransform,
          whiteSpace: cs.whiteSpace,
          paddingLeft: parseFloat(cs.paddingLeft) || 0,
          paddingRight: parseFloat(cs.paddingRight) || 0,
          paddingTop: parseFloat(cs.paddingTop) || 0,
          paddingBottom: parseFloat(cs.paddingBottom) || 0,
        },
        zIndex,
        order: counter.value++
      };
      blocks.push(block);
    }
    return blocks;
  };

  const collectMixedTextBlocks = (slide, slideRect, counter, excludedElements = null) => {
    const blocks = [];
    slide.querySelectorAll("div,section,article").forEach((el) => {
      if (isExcluded(el, excludedElements)) return;
      if (!hasNestedBlocks(el)) return;
      if (el.closest("ul,ol,table")) return;
      const rect = el.getBoundingClientRect();
      if (rect.width < 4 || rect.height < 12) return;
      const runs = collectInlineRuns(el, true);
      const combinedText = runs.map(run => run.text).join("");
      if (!combinedText.trim()) return;
      const directBlocks = Array.from(el.children || []).filter((child) => {
        const tag = child.tagName ? child.tagName.toLowerCase() : "";
        return BLOCK_TAGS.has(tag) && child.innerText && child.innerText.trim().length;
      });
      let y = rect.y - slideRect.y;
      let h = rect.height;
      if (directBlocks.length) {
        const childRects = directBlocks.map((child) => child.getBoundingClientRect());
        const firstTop = Math.min(...childRects.map((childRect) => childRect.top));
        const lastBottom = Math.max(...childRects.map((childRect) => childRect.bottom));
        const topGap = Math.max(0, firstTop - rect.top);
        const bottomGap = Math.max(0, rect.bottom - lastBottom);
        if (topGap >= 8 && topGap >= bottomGap) {
          h = Math.max(0, topGap - 4);
        } else if (bottomGap >= 8) {
          y = Math.min(
            Math.max(lastBottom - slideRect.y + 4, y),
            rect.y - slideRect.y + rect.height - 8
          );
          h = rect.y - slideRect.y + rect.height - y;
        }
      }
      if (h < 8) return;
      const cs = getComputedStyle(el);
      blocks.push({
        kind: "text",
        rect: {
          x: rect.x - slideRect.x,
          y,
          w: rect.width,
          h
        },
        text: combinedText,
        runs,
        styles: {
          fontSizePx: parseFloat(cs.fontSize) || undefined,
          lineHeightPx: parseFloat(cs.lineHeight) || undefined,
          fontWeight: cs.fontWeight,
          fontStyle: cs.fontStyle,
          color: cs.color,
          fontFamily: cs.fontFamily,
          textAlign: cs.textAlign,
          textTransform: cs.textTransform,
          whiteSpace: cs.whiteSpace,
          paddingLeft: parseFloat(cs.paddingLeft) || 0,
          paddingRight: parseFloat(cs.paddingRight) || 0,
          paddingTop: parseFloat(cs.paddingTop) || 0,
          paddingBottom: parseFloat(cs.paddingBottom) || 0,
        },
        zIndex: parseZ(cs.zIndex),
        order: counter.value++
      });
    });
    return blocks;
  };

  const collectBarCharts = (slide, slideRect, counter) => {
    const blocks = [];
    const excludedElements = new Set();
    slide.querySelectorAll("div").forEach((container) => {
      if (container === slide || container.closest("table,svg")) return;
      const containerRect = container.getBoundingClientRect();
      const pushChart = (rowData, labelStyles, labelColumnWidth, extraExcluded) => {
        if (!rowData.length) return;
        const titleEl = Array.from(container.children || []).find((child) => {
          if (!child || !child.tagName) return false;
          return child.classList && child.classList.contains("annotation");
        });
        let title = null;
        if (titleEl) {
          const titleRect = titleEl.getBoundingClientRect();
          const titleStyle = getComputedStyle(titleEl);
          title = {
            text: (titleEl.innerText || "").trim(),
            rect: {
              x: titleRect.x - containerRect.x,
              y: titleRect.y - containerRect.y,
              w: titleRect.width,
              h: titleRect.height,
            },
            styles: {
              fontSizePx: parseFloat(titleStyle.fontSize) || undefined,
              fontWeight: titleStyle.fontWeight,
              fontStyle: titleStyle.fontStyle,
              color: titleStyle.color,
              fontFamily: titleStyle.fontFamily,
              textAlign: titleStyle.textAlign,
            }
          };
        }
        const resolvedMax = Math.max(
          ...rowData.map((row) => row.numericValue || row.ratio * 100 || 0),
          100
        );
        rowData.forEach((row) => {
          if (row.ratio <= 0 && row.numericValue !== null && resolvedMax > 0) {
            row.ratio = Math.max(0, Math.min(1, row.numericValue / resolvedMax));
          }
        });
        blocks.push({
          kind: "bar-chart",
          rect: {
            x: containerRect.x - slideRect.x,
            y: containerRect.y - slideRect.y,
            w: containerRect.width,
            h: containerRect.height
          },
          chart: {
            labelColumnWidth,
            valueColumnWidth: Math.max(...rowData.map((row) => (row.valueRect || {}).w || 0), 0),
            maxValue: resolvedMax,
            labelStyles,
            title,
            rows: rowData
          },
          zIndex: parseZ(getComputedStyle(container).zIndex),
          order: counter.value++
        });
        excludedElements.add(container);
        extraExcluded.forEach((el) => excludedElements.add(el));
      };

      const classBarRows = Array.from(container.querySelectorAll(":scope > .bar-row"));
      if (container.classList.contains("bar-chart") && classBarRows.length >= 2) {
        const labelStyle = getComputedStyle(classBarRows[0].querySelector(".bar-label") || container);
        const rowData = [];
        let labelColumnWidth = 0;
        const elementsToExclude = [container, ...classBarRows];
        classBarRows.forEach((row) => {
          const labelEl = row.querySelector(":scope > .bar-label");
          const trackEl = row.querySelector(":scope > .bar-container");
          const fillEl = trackEl ? trackEl.querySelector(":scope > .bar-fill, :scope > div") : null;
          if (!labelEl || !trackEl || !fillEl) return;
          const rowRect = row.getBoundingClientRect();
          const labelRect = labelEl.getBoundingClientRect();
          const trackRect = trackEl.getBoundingClientRect();
          const fillRect = fillEl.getBoundingClientRect();
          if (trackRect.width < 30 || rowRect.height < 10) return;
          labelColumnWidth = Math.max(labelColumnWidth, labelRect.width);
          const valueText = (fillEl.innerText || "").trim();
          const numericValue = parseNumericText(valueText);
          const ratio = trackRect.width > 0 ? Math.max(0, Math.min(1, fillRect.width / trackRect.width)) : 0;
          const trackStyle = getComputedStyle(trackEl);
          const fillStyle = getComputedStyle(fillEl);
          rowData.push({
            label: (labelEl.innerText || "").trim(),
            numericValue,
            valueText,
            ratio,
            valueInsideFill: true,
            rowRect: {
              x: rowRect.x - containerRect.x,
              y: rowRect.y - containerRect.y,
              w: rowRect.width,
              h: rowRect.height,
            },
            trackRect: {
              x: trackRect.x - containerRect.x,
              y: trackRect.y - containerRect.y,
              w: trackRect.width,
              h: trackRect.height,
            },
            valueRect: {
              x: fillRect.x - containerRect.x,
              y: fillRect.y - containerRect.y,
              w: fillRect.width,
              h: fillRect.height,
            },
            trackStyle: {
              backgroundColor: trackStyle.backgroundColor,
              borderRadius: parseFloat(trackStyle.borderRadius) || 0,
            },
            fillStyle: {
              backgroundColor: fillStyle.backgroundColor,
              borderRadius: parseFloat(fillStyle.borderRadius) || 0,
            },
            valueStyles: {
              fontSizePx: parseFloat(fillStyle.fontSize) || undefined,
              fontWeight: fillStyle.fontWeight,
              fontStyle: fillStyle.fontStyle,
              color: fillStyle.color,
              fontFamily: fillStyle.fontFamily,
              textAlign: "right",
            },
          });
          elementsToExclude.push(labelEl, trackEl, fillEl);
        });
        if (rowData.length >= 2) {
          pushChart(
            rowData,
            {
              fontSizePx: parseFloat(labelStyle.fontSize) || undefined,
              fontWeight: labelStyle.fontWeight,
              fontStyle: labelStyle.fontStyle,
              color: labelStyle.color,
              fontFamily: labelStyle.fontFamily,
              textAlign: labelStyle.textAlign,
            },
            labelColumnWidth,
            elementsToExclude
          );
          return;
        }
      }

      const children = Array.from(container.children || []).filter((child) => child.tagName);
      if (children.length !== 2) return;
      const [labelsEl, barsEl] = children;
      const barRows = Array.from(barsEl.children || []).filter((child) => child.tagName);
      if (barRows.length < 2) return;
      const labels = (labelsEl.innerText || "").split(/\\n+/).map((line) => line.trim()).filter(Boolean);
      if (labels.length !== barRows.length) return;
      const rowData = [];
      let valid = true;

      barRows.forEach((row, index) => {
        const rowChildren = Array.from(row.children || []).filter((child) => child.tagName);
        if (rowChildren.length < 2) {
          valid = false;
          return;
        }
        const trackEl = rowChildren[0];
        const valueEl = rowChildren[rowChildren.length - 1];
        if ((trackEl.tagName || "").toLowerCase() !== "div") {
          valid = false;
          return;
        }
        const fillEl = Array.from(trackEl.children || []).find(
          (child) => (child.tagName || "").toLowerCase() === "div"
        );
        if (!fillEl) {
          valid = false;
          return;
        }

        const rowRect = row.getBoundingClientRect();
        const trackRect = trackEl.getBoundingClientRect();
        const fillRect = fillEl.getBoundingClientRect();
        const valueRect = valueEl.getBoundingClientRect();
        if (rowRect.width < 80 || trackRect.width < 30 || rowRect.height < 10) {
          valid = false;
          return;
        }

        const valueText = (valueEl.innerText || "").trim();
        const numericValue = parseNumericText(valueText);
        let ratio = trackRect.width > 0 ? fillRect.width / trackRect.width : 0;
        ratio = Math.max(0, Math.min(1, ratio));
        if (numericValue === null && ratio <= 0) {
          valid = false;
          return;
        }

        const trackStyle = getComputedStyle(trackEl);
        const fillStyle = getComputedStyle(fillEl);
        const valueStyle = getComputedStyle(valueEl);
        rowData.push({
          label: labels[index],
          numericValue,
          valueText,
          ratio,
          rowRect: {
            x: rowRect.x - container.getBoundingClientRect().x,
            y: rowRect.y - container.getBoundingClientRect().y,
            w: rowRect.width,
            h: rowRect.height,
          },
          trackRect: {
            x: trackRect.x - container.getBoundingClientRect().x,
            y: trackRect.y - container.getBoundingClientRect().y,
            w: trackRect.width,
            h: trackRect.height,
          },
          valueRect: {
            x: valueRect.x - container.getBoundingClientRect().x,
            y: valueRect.y - container.getBoundingClientRect().y,
            w: valueRect.width,
            h: valueRect.height,
          },
          trackStyle: {
            backgroundColor: trackStyle.backgroundColor,
            borderRadius: parseFloat(trackStyle.borderRadius) || 0,
          },
          fillStyle: {
            backgroundColor: fillStyle.backgroundColor,
            borderRadius: parseFloat(fillStyle.borderRadius) || 0,
          },
          valueStyles: {
            fontSizePx: parseFloat(valueStyle.fontSize) || undefined,
            fontWeight: valueStyle.fontWeight,
            fontStyle: valueStyle.fontStyle,
            color: valueStyle.color,
            fontFamily: valueStyle.fontFamily,
            textAlign: valueStyle.textAlign,
          },
        });
      });

      if (!valid || rowData.length !== barRows.length) return;
      const labelsRect = labelsEl.getBoundingClientRect();
      const labelStyle = getComputedStyle(labelsEl);
      const elementsToExclude = [container, labelsEl, barsEl, ...barRows];
      barRows.forEach((row) => {
        Array.from(row.children || []).forEach((child) => {
          elementsToExclude.push(child);
          Array.from(child.children || []).forEach((grandChild) => elementsToExclude.push(grandChild));
        });
      });
      pushChart(
        rowData,
        {
          fontSizePx: parseFloat(labelStyle.fontSize) || undefined,
          fontWeight: labelStyle.fontWeight,
          fontStyle: labelStyle.fontStyle,
          color: labelStyle.color,
          fontFamily: labelStyle.fontFamily,
          textAlign: labelStyle.textAlign,
        },
        labelsRect.width,
        elementsToExclude
      );
    });
    return { blocks, excludedElements };
  };

  const collectLists = (slide, slideRect, counter, excludedElements = null) => {
    const blocks = [];
    slide.querySelectorAll("ul,ol").forEach((list) => {
      if (isExcluded(list, excludedElements)) return;
      if (list.closest("table")) return;
      const rect = list.getBoundingClientRect();
      if (rect.width < 4 || rect.height < 4) return;
      const cs = getComputedStyle(list);
      const itemMeta = [];
      const items = [];
      Array.from(list.querySelectorAll(":scope > li")).forEach((li) => {
        const text = (li.innerText || "").trim();
        if (!text) return;
        const liRect = li.getBoundingClientRect();
        const liStyle = getComputedStyle(li);
        itemMeta.push({
          text,
          rect: {
            x: liRect.x - rect.x,
            y: liRect.y - rect.y,
            w: liRect.width,
            h: liRect.height,
          },
          runs: collectRichTextRuns(li),
          styles: {
            fontSizePx: parseFloat(liStyle.fontSize) || undefined,
            lineHeightPx: parseFloat(liStyle.lineHeight) || undefined,
            fontWeight: liStyle.fontWeight,
            fontStyle: liStyle.fontStyle,
            color: liStyle.color,
            fontFamily: liStyle.fontFamily,
            textAlign: liStyle.textAlign,
            textTransform: liStyle.textTransform,
            whiteSpace: liStyle.whiteSpace,
            paddingLeft: parseFloat(liStyle.paddingLeft) || 0,
            paddingRight: parseFloat(liStyle.paddingRight) || 0,
            paddingTop: parseFloat(liStyle.paddingTop) || 0,
            paddingBottom: parseFloat(liStyle.paddingBottom) || 0,
            marginTop: parseFloat(liStyle.marginTop) || 0,
            marginBottom: parseFloat(liStyle.marginBottom) || 0,
          },
        });
        items.push(text);
      });
      if (!items.length) return;
      blocks.push({
        kind: "list",
        rect: {
          x: rect.x - slideRect.x,
          y: rect.y - slideRect.y,
          w: rect.width,
          h: rect.height
        },
        items,
        ordered: list.tagName.toLowerCase() === "ol",
        styles: {
          fontSizePx: parseFloat(cs.fontSize) || undefined,
          lineHeightPx: parseFloat(cs.lineHeight) || undefined,
          fontWeight: cs.fontWeight,
          fontStyle: cs.fontStyle,
          color: cs.color,
          fontFamily: cs.fontFamily,
          textAlign: cs.textAlign,
          textTransform: cs.textTransform,
          whiteSpace: cs.whiteSpace,
          paddingLeft: parseFloat(cs.paddingLeft) || 0,
          paddingRight: parseFloat(cs.paddingRight) || 0,
          paddingTop: parseFloat(cs.paddingTop) || 0,
          paddingBottom: parseFloat(cs.paddingBottom) || 0,
          listStyleType: cs.listStyleType,
          listStylePosition: cs.listStylePosition,
        },
        itemMeta,
        zIndex: parseZ(cs.zIndex),
        order: counter.value++
      });
    });
    return blocks;
  };

  const collectTables = (slide, slideRect, counter, excludedElements = null) => {
    const blocks = [];
    slide.querySelectorAll("table").forEach((tbl) => {
      if (isExcluded(tbl, excludedElements)) return;
      const rect = tbl.getBoundingClientRect();
      if (rect.width < 4 || rect.height < 4) return;
      const rows = [];
      const rowCells = [];
      const columnWidths = [];
      const rowHeights = [];
      tbl.querySelectorAll("tr").forEach((tr) => {
        const row = [];
        const cells = [];
        const trRect = tr.getBoundingClientRect();
        rowHeights.push(trRect.height);
        let colIndex = 0;
        tr.querySelectorAll("th,td").forEach((cell) => {
          const text = (cell.innerText || "").trim();
          const csCell = getComputedStyle(cell);
          const cellRect = cell.getBoundingClientRect();
          const colSpan = parseInt(cell.getAttribute("colspan") || "1", 10) || 1;
          const avgWidth = colSpan > 0 ? cellRect.width / colSpan : cellRect.width;
          for (let offset = 0; offset < colSpan; offset++) {
            const targetIndex = colIndex + offset;
            columnWidths[targetIndex] = Math.max(columnWidths[targetIndex] || 0, avgWidth);
          }
          colIndex += colSpan;
          row.push(text);
          cells.push({
            text,
            runs: collectRichTextRuns(cell),
            isHeader: (cell.tagName || "").toLowerCase() === "th",
            colSpan,
            rowSpan: parseInt(cell.getAttribute("rowspan") || "1", 10) || 1,
            rect: {
              w: cellRect.width,
              h: cellRect.height
            },
            styles: {
              backgroundColor: csCell.backgroundColor,
              borderColor: csCell.borderColor,
              borderWidth: parseFloat(csCell.borderWidth) || 0,
              color: csCell.color,
              fontSizePx: parseFloat(csCell.fontSize) || undefined,
              lineHeightPx: parseFloat(csCell.lineHeight) || undefined,
              fontWeight: csCell.fontWeight,
              fontStyle: csCell.fontStyle,
              fontFamily: csCell.fontFamily,
              textAlign: csCell.textAlign,
              verticalAlign: csCell.verticalAlign,
              paddingLeft: parseFloat(csCell.paddingLeft) || 0,
              paddingRight: parseFloat(csCell.paddingRight) || 0,
              paddingTop: parseFloat(csCell.paddingTop) || 0,
              paddingBottom: parseFloat(csCell.paddingBottom) || 0
            }
          });
        });
        if (row.length) {
          rows.push(row);
          rowCells.push(cells);
        }
      });
      if (!rows.length) return;
      const cs = getComputedStyle(tbl);
      blocks.push({
        kind: "table",
        rect: {
          x: rect.x - slideRect.x,
          y: rect.y - slideRect.y,
          w: rect.width,
          h: rect.height
        },
        rows,
        rowCells,
        columnWidths,
        rowHeights,
        styles: {
          fontSizePx: parseFloat(cs.fontSize) || undefined,
          color: cs.color,
          fontFamily: cs.fontFamily,
          textAlign: cs.textAlign,
        },
        zIndex: parseZ(cs.zIndex),
        order: counter.value++
      });
    });
    return blocks;
  };

  const collectImages = (slide, slideRect, counter, excludedElements = null) => {
    const blocks = [];
    slide.querySelectorAll("img").forEach((img) => {
      if (isExcluded(img, excludedElements)) return;
      const rect = img.getBoundingClientRect();
      if (rect.width < 4 || rect.height < 4) return;
      const cs = getComputedStyle(img);
      let dataUrl = null;
      try {
        if (img.complete && img.naturalWidth > 0 && img.naturalHeight > 0) {
          const maxDim = 2048;
          const scale = Math.min(1, maxDim / Math.max(img.naturalWidth, img.naturalHeight));
          const canvas = document.createElement("canvas");
          canvas.width = Math.max(1, Math.round(img.naturalWidth * scale));
          canvas.height = Math.max(1, Math.round(img.naturalHeight * scale));
          const ctx = canvas.getContext("2d");
          if (ctx) {
            ctx.drawImage(img, 0, 0, canvas.width, canvas.height);
            dataUrl = canvas.toDataURL("image/png");
          }
        }
      } catch (err) {
        dataUrl = null;
      }
      blocks.push({
        kind: "image",
        rect: {
          x: rect.x - slideRect.x,
          y: rect.y - slideRect.y,
          w: rect.width,
          h: rect.height
        },
        src: img.currentSrc || img.src || "",
        alt: img.alt || "",
        dataUrl,
        naturalWidth: img.naturalWidth || undefined,
        naturalHeight: img.naturalHeight || undefined,
        zIndex: parseZ(cs.zIndex),
        order: counter.value++
      });
    });
    return blocks;
  };

  const collectShapes = (slide, slideRect, counter, excludedElements = null) => {
    const blocks = [];
    slide.querySelectorAll("*").forEach((el) => {
      if (el === slide) return;
      if (isExcluded(el, excludedElements)) return;
      const tag = el.tagName ? el.tagName.toLowerCase() : "";
      if (["table", "thead", "tbody", "tfoot", "tr", "td", "th", "svg"].includes(tag)) {
        return;
      }
      if (el.closest("table, svg")) return;
      const rect = el.getBoundingClientRect();
      if (rect.width < 6 || rect.height < 6) return;
      const cs = getComputedStyle(el);
      const bgImage = cs.backgroundImage || "";
      const approxGradientFill = approximateLinearGradientColor(bgImage);
      const bg = cs.backgroundColor && !/^rgba?\(0,\s*0,\s*0,\s*0\)/i.test(cs.backgroundColor)
        ? cs.backgroundColor
        : approxGradientFill;
      const borderWidth = parseFloat(cs.borderWidth) || 0;
      const borderColor = cs.borderColor;
      const borderTopWidth = parseFloat(cs.borderTopWidth) || 0;
      const borderRightWidth = parseFloat(cs.borderRightWidth) || 0;
      const borderBottomWidth = parseFloat(cs.borderBottomWidth) || 0;
      const borderLeftWidth = parseFloat(cs.borderLeftWidth) || 0;
      const radius = parseFloat(cs.borderRadius) || 0;
      const hasFill = Boolean(bg && !/^rgba?\(0,\s*0,\s*0,\s*0\)/i.test(bg));
      const hasBorder =
        (borderWidth > 0 && borderColor && borderColor !== "rgba(0, 0, 0, 0)") ||
        borderTopWidth > 0 ||
        borderRightWidth > 0 ||
        borderBottomWidth > 0 ||
        borderLeftWidth > 0;
      if (!hasFill && !hasBorder) return;
      blocks.push({
        kind: "shape",
        rect: {
          x: rect.x - slideRect.x,
          y: rect.y - slideRect.y,
          w: rect.width,
          h: rect.height
        },
        shape: {
          fill: bg,
          borderColor: borderColor,
          borderWidth,
          borderTopWidth,
          borderRightWidth,
          borderBottomWidth,
          borderLeftWidth,
          borderTopColor: cs.borderTopColor,
          borderRightColor: cs.borderRightColor,
          borderBottomColor: cs.borderBottomColor,
          borderLeftColor: cs.borderLeftColor,
          borderRadius: radius,
          isEllipse:
            rect.width > 20 &&
            rect.height > 20 &&
            Math.abs(rect.width - rect.height) <= Math.max(rect.width, rect.height) * 0.15 &&
            ((cs.borderRadius || "").includes("%") || radius >= Math.min(rect.width, rect.height) / 2 - 2)
        },
        zIndex: parseZ(cs.zIndex),
        order: counter.value++
      });
    });
    return blocks;
  };

  const splitGradientParts = (body) => {
    const parts = [];
    let depth = 0;
    let current = "";
    for (let i = 0; i < body.length; i++) {
      const ch = body[i];
      if (ch === "(") {
        depth++;
        current += ch;
      } else if (ch === ")") {
        depth = Math.max(0, depth - 1);
        current += ch;
      } else if (ch === "," && depth === 0) {
        if (current.trim().length) {
          parts.push(current.trim());
        }
        current = "";
      } else {
        current += ch;
      }
    }
    if (current.trim().length) {
      parts.push(current.trim());
    }
    return parts;
  };

  const parseConicGradient = (bgImage) => {
    const start = bgImage.indexOf("(");
    const end = bgImage.lastIndexOf(")");
    if (start < 0 || end <= start) return null;
    const body = bgImage.slice(start + 1, end);
    const parts = splitGradientParts(body);
    const segments = [];
    let prevEnd = 0;
    for (const part of parts) {
      const trimmed = part.trim();
      if (!trimmed || trimmed.startsWith("from ") || trimmed.startsWith("at ")) {
        continue;
      }
      const tokens = trimmed.split(/\s+/);
      const colorTokens = [];
      const positionTokens = [];
      let collectingColor = true;
      for (let i = 0; i < tokens.length; i++) {
        let token = tokens[i];
        if (!token) continue;
        if (collectingColor) {
          colorTokens.push(token);
          if (token.includes("(") && !token.includes(")")) {
            while (i + 1 < tokens.length && !tokens[i].includes(")")) {
              i++;
              colorTokens.push(tokens[i]);
              if (tokens[i].includes(")")) break;
            }
          }
          collectingColor = false;
        } else {
          positionTokens.push(token);
        }
      }
      const color = colorTokens.join(" ");
      if (!color) continue;
      const toPercent = (value) => {
        if (!value) return null;
        if (value.includes("%")) {
          return parseFloat(value);
        }
        if (value.includes("deg")) {
          const deg = parseFloat(value);
          return (deg / 360) * 100;
        }
        return parseFloat(value);
      };
      let startPct = prevEnd;
      let endPct = null;
      if (positionTokens.length >= 2) {
        startPct = toPercent(positionTokens[0]);
        endPct = toPercent(positionTokens[1]);
      } else if (positionTokens.length === 1) {
        endPct = toPercent(positionTokens[0]);
      }
      if (endPct === null || startPct === null) continue;
      if (endPct < startPct) {
        const temp = startPct;
        startPct = endPct;
        endPct = temp;
      }
      segments.push({
        color,
        startDeg: (startPct / 100) * 360,
        endDeg: (endPct / 100) * 360
      });
      prevEnd = endPct;
    }
    return segments.length ? segments : null;
  };

  const collectConicGradients = (slide, slideRect, counter, excludedElements = null) => {
    const blocks = [];
    slide.querySelectorAll("*").forEach((el) => {
      if (isExcluded(el, excludedElements)) return;
      const rect = el.getBoundingClientRect();
      if (rect.width < 20 || rect.height < 20) return;
      const cs = getComputedStyle(el);
      const bgImage = cs.backgroundImage || "";
      if (!bgImage.includes("conic-gradient")) return;
      const segments = parseConicGradient(bgImage);
      if (!segments || !segments.length) return;
      blocks.push({
        kind: "conic-gradient",
        rect: {
          x: rect.x - slideRect.x,
          y: rect.y - slideRect.y,
          w: rect.width,
          h: rect.height
        },
        gradient: {
          cx: rect.x - slideRect.x + rect.width / 2,
          cy: rect.y - slideRect.y + rect.height / 2,
          radius: Math.min(rect.width, rect.height) / 2,
          segments
        },
        zIndex: parseZ(cs.zIndex),
        order: counter.value++
      });
    });
    return blocks;
  };

  const collectSvgElements = (slide, slideRect, counter, excludedElements = null) => {
    const result = [];
    const parseLength = (value, fallback = 0) => {
      if (value === undefined || value === null) return fallback;
      if (typeof value === "number") return value;
      const num = parseFloat(value);
      return Number.isFinite(num) ? num : fallback;
    };
    slide.querySelectorAll("svg").forEach((svg) => {
      if (isExcluded(svg, excludedElements)) return;
      const svgRect = svg.getBoundingClientRect();
      if (svgRect.width < 2 || svgRect.height < 2) return;
      const viewBox = svg.viewBox && svg.viewBox.baseVal ? svg.viewBox.baseVal : null;
      const baseWidth = viewBox && viewBox.width ? viewBox.width : parseLength(svg.getAttribute("width"), svgRect.width);
      const baseHeight = viewBox && viewBox.height ? viewBox.height : parseLength(svg.getAttribute("height"), svgRect.height);
      const scaleX = baseWidth ? svgRect.width / baseWidth : 1;
      const scaleY = baseHeight ? svgRect.height / baseHeight : 1;
      const offsetX = svgRect.x - slideRect.x;
      const offsetY = svgRect.y - slideRect.y;

      const toGlobalPoints = (points) => {
        const arr = [];
        for (let i = 0; i < points.length; i++) {
          const pt = points[i];
          arr.push({
            x: offsetX + pt.x * scaleX,
            y: offsetY + pt.y * scaleY,
          });
        }
        return arr;
      };

      const svgStyle = getComputedStyle(svg);
      const baseZ = parseZ(svgStyle.zIndex);

      svg.querySelectorAll("polyline,polygon,line").forEach((el) => {
        const cs = getComputedStyle(el);
        const stroke = cs.stroke && cs.stroke !== "none" ? cs.stroke : (el.getAttribute("stroke") || null);
        const strokeWidth = parseLength(cs.strokeWidth, parseLength(el.getAttribute("stroke-width"), 1));
        const fill = cs.fill && cs.fill !== "none" ? cs.fill : (el.getAttribute("fill") || null);
        let pts = [];
        if (el.tagName.toLowerCase() === "line") {
          const x1 = parseLength(el.getAttribute("x1"));
          const y1 = parseLength(el.getAttribute("y1"));
          const x2 = parseLength(el.getAttribute("x2"));
          const y2 = parseLength(el.getAttribute("y2"));
          pts = [
            { x: offsetX + x1 * scaleX, y: offsetY + y1 * scaleY },
            { x: offsetX + x2 * scaleX, y: offsetY + y2 * scaleY },
          ];
        } else {
          if (el.points && el.points.length) {
            pts = toGlobalPoints(el.points);
          } else {
            const attr = el.getAttribute("points") || "";
            const pairRegex = /(-?\\d+(?:\\.\\d+)?),(-?\\d+(?:\\.\\d+)?)/g;
            let match;
            while ((match = pairRegex.exec(attr)) !== null) {
              const x = parseFloat(match[1]);
              const y = parseFloat(match[2]);
              if (Number.isFinite(x) && Number.isFinite(y)) {
                pts.push({ x: offsetX + x * scaleX, y: offsetY + y * scaleY });
              }
            }
          }
        }
        if (pts.length >= 2) {
          result.push({
            kind: "svg-polyline",
            polyline: {
              points: pts,
              stroke,
              strokeWidth,
              fill,
              closed: el.tagName.toLowerCase() === "polygon",
            },
            rect: {
              x: offsetX,
              y: offsetY,
              w: svgRect.width,
              h: svgRect.height,
            },
            zIndex: baseZ,
            order: counter.value++
          });
        }
      });

      svg.querySelectorAll("circle").forEach((el) => {
        const cs = getComputedStyle(el);
        const cx = parseLength(el.getAttribute("cx"));
        const cy = parseLength(el.getAttribute("cy"));
        const r = parseLength(el.getAttribute("r"));
        if (!r) return;
        const stroke = cs.stroke && cs.stroke !== "none" ? cs.stroke : (el.getAttribute("stroke") || null);
        const strokeWidth = parseLength(cs.strokeWidth, parseLength(el.getAttribute("stroke-width"), 1));
        const fill = cs.fill && cs.fill !== "none" ? cs.fill : (el.getAttribute("fill") || null);
        result.push({
          kind: "svg-circle",
          circle: {
            cx: offsetX + cx * scaleX,
            cy: offsetY + cy * scaleY,
            r: r * Math.max(scaleX, scaleY),
            stroke,
            strokeWidth,
            fill,
          },
          rect: {
            x: offsetX,
            y: offsetY,
            w: svgRect.width,
            h: svgRect.height,
          },
          zIndex: baseZ,
          order: counter.value++
        });
      });

      svg.querySelectorAll("ellipse").forEach((el) => {
        const cs = getComputedStyle(el);
        const cx = parseLength(el.getAttribute("cx"));
        const cy = parseLength(el.getAttribute("cy"));
        const rx = parseLength(el.getAttribute("rx"));
        const ry = parseLength(el.getAttribute("ry"));
        if (!rx || !ry) return;
        const stroke = cs.stroke && cs.stroke !== "none" ? cs.stroke : (el.getAttribute("stroke") || null);
        const strokeWidth = parseLength(cs.strokeWidth, parseLength(el.getAttribute("stroke-width"), 1));
        const fill = cs.fill && cs.fill !== "none" ? cs.fill : (el.getAttribute("fill") || null);
        result.push({
          kind: "svg-ellipse",
          ellipse: {
            cx: offsetX + cx * scaleX,
            cy: offsetY + cy * scaleY,
            rx: rx * scaleX,
            ry: ry * scaleY,
            stroke,
            strokeWidth,
            fill,
          },
          rect: {
            x: offsetX,
            y: offsetY,
            w: svgRect.width,
            h: svgRect.height,
          },
          zIndex: baseZ,
          order: counter.value++
        });
      });
    });
    return result;
  };

  return targets.map((slide) => {
    const rect = slide.getBoundingClientRect();
    const cs = getComputedStyle(slide);
    let background = cs.backgroundColor;
    if (!background || background === "rgba(0, 0, 0, 0)") {
      const bodyStyle = getComputedStyle(document.body);
      background = bodyStyle.backgroundColor || background;
    }
    const counter = { value: 0 };
    const barCharts = collectBarCharts(slide, rect, counter);
    const blocks = [
      ...barCharts.blocks,
      ...collectShapes(slide, rect, counter, barCharts.excludedElements),
      ...collectTextBlocks(slide, rect, counter, barCharts.excludedElements),
      ...collectMixedTextBlocks(slide, rect, counter, barCharts.excludedElements),
      ...collectLists(slide, rect, counter, barCharts.excludedElements),
      ...collectTables(slide, rect, counter, barCharts.excludedElements),
      ...collectImages(slide, rect, counter, barCharts.excludedElements),
      ...collectSvgElements(slide, rect, counter, barCharts.excludedElements),
      ...collectConicGradients(slide, rect, counter, barCharts.excludedElements)
    ];
    blocks.sort((a, b) => a.order - b.order);
    return {
      viewport: { w: window.innerWidth, h: window.innerHeight },
      rect: { x: rect.x, y: rect.y, w: rect.width, h: rect.height },
      background,
      blocks
    };
  });
}
"""

def file_or_url_to_uri(path_or_url: str) -> str:
    p = str(path_or_url)
    if p.startswith("http://") or p.startswith("https://"):
        return p
    return Path(p).resolve().as_uri()


@dataclass
class LayoutBox:
    left: Optional[float] = None
    top: Optional[float] = None
    width: Optional[float] = None
    height: Optional[float] = None


@dataclass
class TextRun:
    text: str
    font_size: Optional[float] = None
    bold: Optional[bool] = None
    italic: Optional[bool] = None
    color: Optional[str] = None
    font_family: Optional[str] = None


@dataclass
class TableCell:
    text: str
    background_color: Optional[str] = None
    border_color: Optional[str] = None
    border_width: Optional[float] = None
    text_style: Dict[str, Any] = field(default_factory=dict)
    vertical_align: Optional[str] = None
    runs: List[TextRun] = field(default_factory=list)
    is_header: bool = False
    padding_left: Optional[float] = None
    padding_right: Optional[float] = None
    padding_top: Optional[float] = None
    padding_bottom: Optional[float] = None


@dataclass
class Block:
    kind: str  # text, list, table, image, shape, polyline, circle, ellipse, bar-chart
    text: str = ""
    runs: List[TextRun] = field(default_factory=list)
    items: List[str] = field(default_factory=list)
    numbered: bool = False
    table: List[List[str]] = field(default_factory=list)
    table_cells: List[List[TableCell]] = field(default_factory=list)
    image_path: Optional[Path] = None
    image_alt: str = ""
    shape_style: Dict[str, Any] = field(default_factory=dict)
    text_style: Dict[str, Any] = field(default_factory=dict)
    vector_data: Dict[str, Any] = field(default_factory=dict)
    layout: LayoutBox = field(default_factory=LayoutBox)
    position: Optional[str] = None  # static / relative / absolute / fixed
    z_index: int = 0
    order: int = 0


@dataclass
class SlideModel:
    title: Optional[str]
    background_color: Optional[str]
    blocks: List[Block] = field(default_factory=list)
    layout_constraints: List["LayoutConstraint"] = field(default_factory=list)
    canvas_width: float = SLIDE_REF_WIDTH
    canvas_height: float = SLIDE_REF_HEIGHT
    scale: Optional[float] = None
    scale_x: Optional[float] = None
    scale_y: Optional[float] = None
    offset_x: float = 0.0
    offset_y: float = 0.0
    transform_mode: str = "contain"
    raster_image_data_url: Optional[str] = None
    raster_reason: Optional[str] = None


@dataclass
class LayoutSlot:
    element_id: int
    blocks: List[Block] = field(default_factory=list)


@dataclass
class LayoutConstraint:
    kind: str  # grid or flex
    parent_style: Dict[str, str]
    parent_tag: Tag
    depth: int
    slots: List[LayoutSlot] = field(default_factory=list)
    _slot_map: Dict[int, LayoutSlot] = field(default_factory=dict, repr=False)


def compute_dom_depth(element: Tag) -> int:
    depth = 0
    current = element
    while isinstance(current, Tag):
        parent = current.parent
        if not isinstance(parent, Tag):
            break
        depth += 1
        current = parent
    return depth


def find_direct_child_for_parent(element: Tag, ancestor: Tag) -> Optional[Tag]:
    child = element
    parent = element.parent
    while parent and parent is not ancestor:
        if not isinstance(parent, Tag):
            return None
        child = parent
        parent = parent.parent
    return child if parent is ancestor else None


def register_layout_constraints(
    element: Tag,
    block: Block,
    resolver: "StyleResolver",
    constraints: Dict[int, LayoutConstraint],
) -> None:
    parent = element.parent
    while isinstance(parent, Tag):
        style = resolver.get_style(parent)
        display = style.get("display", "").lower()
        if display in {"grid", "flex"}:
            key = id(parent)
            constraint = constraints.get(key)
            if not constraint:
                constraint = LayoutConstraint(
                    kind=display,
                    parent_style=dict(style),
                    parent_tag=parent,
                    depth=resolver.dom_depth(parent),
                )
                constraints[key] = constraint
            direct_child = find_direct_child_for_parent(element, parent)
            if direct_child is not None:
                slot_id = id(direct_child)
                slot = constraint._slot_map.get(slot_id)
                if not slot:
                    slot = LayoutSlot(element_id=slot_id)
                    constraint._slot_map[slot_id] = slot
                    constraint.slots.append(slot)
                slot.blocks.append(block)
        parent = parent.parent if isinstance(parent.parent, Tag) else None


def parse_gap_values(style: Dict[str, str]) -> Tuple[float, float]:
    row_gap = parse_length(style.get("row-gap"))
    col_gap = parse_length(style.get("column-gap"))
    gap_val = style.get("gap")
    if gap_val:
        parts = gap_val.split()
        if len(parts) == 1:
            gap = parse_length(parts[0])
            if gap is not None:
                if row_gap is None:
                    row_gap = gap
                if col_gap is None:
                    col_gap = gap
        elif len(parts) >= 2:
            first = parse_length(parts[0])
            second = parse_length(parts[1])
            if row_gap is None:
                row_gap = first
            if col_gap is None:
                col_gap = second
    if row_gap is None:
        row_gap = FLOW_GAP
    if col_gap is None:
        col_gap = FLOW_GAP
    return row_gap, col_gap


def expand_grid_template_tokens(template: str) -> List[str]:
    tokens: List[str] = []
    if not template:
        return tokens
    buffer = ""
    depth = 0
    for ch in template:
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth = max(0, depth - 1)
        if ch.isspace() and depth == 0:
            if buffer:
                tokens.append(buffer.strip())
                buffer = ""
        else:
            buffer += ch
    if buffer:
        tokens.append(buffer.strip())
    expanded: List[str] = []
    for token in tokens:
        if not token:
            continue
        match = GRID_REPEAT_RE.match(token)
        if match:
            count = int(match.group(1))
            inner = match.group(2).strip()
            for _ in range(count):
                expanded.append(inner)
        else:
            expanded.append(token)
    return expanded


def compute_grid_column_widths(
    container_width: float,
    template: Optional[str],
    column_gap: float,
    block_count: int,
) -> List[float]:
    tokens = expand_grid_template_tokens(template or "")
    if not tokens:
        count = max(1, min(block_count, 4))
        available = max(container_width - column_gap * (count - 1), 100.0)
        width = available / count
        return [width for _ in range(count)]
    column_defs: List[Tuple[str, float]] = []
    px_total = 0.0
    fr_total = 0.0
    for token in tokens:
        norm = token.lower()
        if norm.endswith("fr"):
            try:
                value = float(norm[:-2] or "1")
            except ValueError:
                value = 1.0
            fr_total += value
            column_defs.append(("fr", value))
        else:
            length = parse_length(token)
            if length is None:
                fr_total += 1.0
                column_defs.append(("fr", 1.0))
            else:
                px_total += length
                column_defs.append(("px", length))
    count = len(column_defs)
    available = max(container_width - column_gap * (count - 1), 50.0)
    widths: List[float] = []
    if px_total > available and px_total > 0:
        scale = available / px_total
        column_defs = [
            (kind, value * scale) if kind == "px" else (kind, value) for kind, value in column_defs
        ]
        px_total = available
    remaining = max(available - px_total, 0.0)
    for kind, value in column_defs:
        if kind == "px":
            widths.append(value)
        else:
            portion = remaining / fr_total if fr_total else (remaining / count if count else remaining)
            widths.append(portion * value if fr_total else portion)
    return widths


def compute_slot_metrics(slot: LayoutSlot) -> Tuple[float, float, float, float]:
    min_left = None
    min_top = None
    max_right = None
    max_bottom = None
    for block in slot.blocks:
        left = block.layout.left if block.layout.left is not None else DEFAULT_PADDING_X
        top = block.layout.top if block.layout.top is not None else DEFAULT_PADDING_Y
        width = block.layout.width if block.layout.width is not None else SLIDE_REF_WIDTH - 2 * DEFAULT_PADDING_X
        height = block.layout.height if block.layout.height is not None else estimate_block_height(block)
        min_left = left if min_left is None else min(min_left, left)
        min_top = top if min_top is None else min(min_top, top)
        right = left + width
        bottom = top + height
        max_right = right if max_right is None else max(max_right, right)
        max_bottom = bottom if max_bottom is None else max(max_bottom, bottom)
    min_left = min_left if min_left is not None else DEFAULT_PADDING_X
    min_top = min_top if min_top is not None else DEFAULT_PADDING_Y
    width_span = (max_right - min_left) if max_right is not None else SLIDE_REF_WIDTH - 2 * DEFAULT_PADDING_X
    height_span = (max_bottom - min_top) if max_bottom is not None else DEFAULT_BLOCK_HEIGHT
    return float(min_left), float(min_top), float(max(width_span, 10.0)), float(max(height_span, 10.0))


def compute_container_box(constraint: LayoutConstraint) -> Tuple[float, float, float]:
    layout = extract_layout(constraint.parent_style, constraint.parent_tag)
    block_lefts = [
        block.layout.left
        for slot in constraint.slots
        for block in slot.blocks
        if block.layout.left is not None
    ]
    block_rights = [
        (block.layout.left + block.layout.width)
        for slot in constraint.slots
        for block in slot.blocks
        if block.layout.left is not None and block.layout.width is not None
    ]
    block_tops = [
        block.layout.top
        for slot in constraint.slots
        for block in slot.blocks
        if block.layout.top is not None
    ]
    left = layout.left if layout.left is not None else (min(block_lefts) if block_lefts else DEFAULT_PADDING_X)
    top = layout.top if layout.top is not None else (min(block_tops) if block_tops else DEFAULT_PADDING_Y)
    if layout.width is not None:
        width = layout.width
    elif block_rights:
        width = max(block_rights) - left
    else:
        width = SLIDE_REF_WIDTH - 2 * DEFAULT_PADDING_X
    return float(left), float(top), float(max(width, 100.0))


def px_to_pt(px: float) -> float:
    return float(px) * 0.75


@lru_cache(maxsize=1024)
def css_color_to_rgb_tuple(color_str: Optional[str]) -> Optional[Tuple[int, int, int]]:
    if not color_str:
        return None
    s = color_str.strip()
    if not s:
        return None
    try:
        if s.startswith("#"):
            if len(s) == 4:  # #RGB
                r = int(s[1], 16) * 17
                g = int(s[2], 16) * 17
                b = int(s[3], 16) * 17
                return (r, g, b)
            if len(s) == 7:
                return (int(s[1:3], 16), int(s[3:5], 16), int(s[5:7], 16))
        if s.lower().startswith("rgb"):
            nums = s[s.find("(") + 1 : s.find(")")].split(",")
            r, g, b = [int(float(v.strip())) for v in nums[:3]]
            alpha = float(nums[3].strip()) if len(nums) >= 4 else 1.0
            if alpha < 1.0:
                r = int(round(r * alpha + 255 * (1.0 - alpha)))
                g = int(round(g * alpha + 255 * (1.0 - alpha)))
                b = int(round(b * alpha + 255 * (1.0 - alpha)))
            return (r, g, b)
    except Exception:
        return None
    keyword = CSS_COLOR_KEYWORDS.get(s.lower())
    if keyword:
        return keyword
    _, image_color = ensure_pillow()
    if image_color:
        try:
            return tuple(image_color.getrgb(s))
        except Exception:
            return None
    return None


@lru_cache(maxsize=1024)
def css_is_transparent(color_str: Optional[str]) -> bool:
    if not color_str:
        return True
    s = color_str.strip().lower()
    if not s:
        return True
    if s == "transparent":
        return True
    if s.startswith("rgba"):
        try:
            nums = s[s.find("(") + 1 : s.find(")")].split(",")
            if len(nums) >= 4:
                alpha = float(nums[3])
                return alpha <= 0.0
        except Exception:
            return False
    return False


def parse_length(value: Optional[str], reference: Optional[float] = None) -> Optional[float]:
    if value is None:
        return None
    s = value.strip()
    if not s or s in {"auto", "initial", "inherit"}:
        return None
    if s.endswith("px"):
        return float(s[:-2])
    if s.endswith("pt"):
        return float(s[:-2]) * (96.0 / 72.0)
    if s.endswith("in"):
        return float(s[:-2]) * 96.0
    if s.endswith("cm"):
        return float(s[:-2]) * (96.0 / 2.54)
    if s.endswith("mm"):
        return float(s[:-2]) * (96.0 / 25.4)
    if s.endswith("%") and reference is not None:
        try:
            return float(s[:-1]) / 100.0 * reference
        except ValueError:
            return None
    if s.endswith("em") or s.endswith("rem"):
        try:
            return float(s[:-2]) * DEFAULT_FONT_SIZE
        except ValueError:
            return None
    try:
        return float(s)
    except ValueError:
        return None


def parse_line_height(value: Optional[str], font_size: Optional[float] = None) -> Optional[float]:
    if value is None:
        return None
    s = value.strip().lower()
    if not s or s == "normal":
        return None
    if s.endswith("px"):
        try:
            return float(s[:-2])
        except ValueError:
            return None
    if s.endswith("%") and font_size:
        try:
            return float(s[:-1]) / 100.0 * font_size
        except ValueError:
            return None
    try:
        num = float(s)
        if font_size:
            return num * font_size if num <= 10 else num
        return num if num > 0 else None
    except ValueError:
        return None


def parse_font_size(value: Optional[str], tag: Optional[str] = None) -> float:
    size = parse_length(value)
    if size:
        return size
    if tag and tag.lower() in TAG_FONT_SIZE:
        return TAG_FONT_SIZE[tag.lower()]
    return DEFAULT_FONT_SIZE


def split_font_family_list(family: Optional[str]) -> List[str]:
    if not family:
        return []
    parts = []
    for part in family.split(","):
        normalized = part.strip().strip("'\"")
        if normalized:
            parts.append(normalized)
    return parts


def resolve_font_family_name(family: Optional[str]) -> Optional[str]:
    for name in split_font_family_list(family):
        lowered = name.lower()
        if lowered in GENERIC_FONT_MAPPING:
            continue
        return name
    for name in split_font_family_list(family):
        mapped = GENERIC_FONT_MAPPING.get(name.lower())
        if mapped:
            return mapped
    return None


def normalize_whitespace(text: str) -> str:
    return WHITESPACE_RE.sub(" ", text).strip()


def apply_text_transform(text: str, transform: Optional[str]) -> str:
    if not text:
        return text
    normalized = (transform or "").strip().lower()
    if normalized in {"", "none"}:
        return text
    if normalized == "uppercase":
        return text.upper()
    if normalized == "lowercase":
        return text.lower()
    if normalized == "capitalize":
        return re.sub(r"(^|[\s\u3000])(\S)", lambda m: f"{m.group(1)}{m.group(2).upper()}", text)
    return text


@lru_cache(maxsize=2048)
def _parse_declarations_cached(text: str) -> Tuple[Tuple[str, str], ...]:
    result: List[Tuple[str, str]] = []
    for part in text.split(";"):
        if ":" not in part:
            continue
        name, val = part.split(":", 1)
        name = name.strip().lower()
        if not name:
            continue
        result.append((name, val.strip()))
    return tuple(result)


def parse_declarations(text: str) -> Dict[str, str]:
    return dict(_parse_declarations_cached(text or ""))


@dataclass
class StyleRule:
    tag: Optional[str]
    element_id: Optional[str]
    classes: Tuple[str, ...]
    declarations: Dict[str, str]
    order: int = 0
    class_lookup: frozenset[str] = field(init=False, repr=False)

    def __post_init__(self) -> None:
        self.class_lookup = frozenset(self.classes)

    def matches(self, tag: str, class_set: frozenset[str], element_id: Optional[str]) -> bool:
        if self.tag and self.tag != "*" and self.tag != tag:
            return False
        if self.element_id and self.element_id != element_id:
            return False
        if self.class_lookup and not self.class_lookup.issubset(class_set):
            return False
        return True


def parse_simple_selector(selector: str) -> Optional[StyleRule]:
    selector = selector.strip()
    if not selector or " " in selector or ">" in selector:
        return None
    tokens = SIMPLE_SELECTOR_TOKEN_RE.findall(selector)
    if not tokens:
        return None
    tag: Optional[str] = None
    el_id: Optional[str] = None
    classes: List[str] = []
    for token in tokens:
        if token == "*":
            tag = "*"
        elif token.startswith("#"):
            el_id = token[1:]
        elif token.startswith("."):
            classes.append(token[1:])
        else:
            tag = token.lower()
    return StyleRule(tag=tag, element_id=el_id, classes=tuple(classes), declarations={})


class StyleResolver:
    """Very small CSS resolver (supports tag/id/class selectors without combinators)."""

    def __init__(self, soup: BeautifulSoup):
        self.rules: List[StyleRule] = []
        self.rules_by_tag: Dict[str, List[StyleRule]] = {}
        self.rules_by_class: Dict[str, List[StyleRule]] = {}
        self.rules_by_id: Dict[str, List[StyleRule]] = {}
        self.universal_rules: List[StyleRule] = []
        self._style_cache: Dict[int, Dict[str, str]] = {}
        self._depth_cache: Dict[int, int] = {}
        self._rule_order = 0
        for style_tag in soup.select("style"):
            self._consume_stylesheet(style_tag.string or "")

    def _index_rule(self, rule: StyleRule) -> None:
        if not rule.tag or rule.tag == "*":
            self.universal_rules.append(rule)
        else:
            self.rules_by_tag.setdefault(rule.tag, []).append(rule)
        if rule.element_id:
            self.rules_by_id.setdefault(rule.element_id, []).append(rule)
        for cls in rule.classes:
            self.rules_by_class.setdefault(cls, []).append(rule)

    def _consume_stylesheet(self, css_text: str) -> None:
        cleaned = CSS_COMMENT_RE.sub("", css_text)
        for raw_rule in cleaned.split("}"):
            if "{" not in raw_rule:
                continue
            selector_text, body = raw_rule.split("{", 1)
            declarations = parse_declarations(body)
            if not declarations:
                continue
            for selector in selector_text.split(","):
                rule = parse_simple_selector(selector)
                if not rule:
                    continue
                rule.declarations = declarations.copy()
                rule.order = self._rule_order
                self._rule_order += 1
                self.rules.append(rule)
                self._index_rule(rule)

    def dom_depth(self, element: Tag) -> int:
        cache_key = id(element)
        cached = self._depth_cache.get(cache_key)
        if cached is not None:
            return cached
        depth = 0
        current = element
        while isinstance(current, Tag):
            parent = current.parent
            if not isinstance(parent, Tag):
                break
            depth += 1
            current = parent
        self._depth_cache[cache_key] = depth
        return depth

    def get_style(self, element: Tag) -> Dict[str, str]:
        cache_key = id(element)
        cached = self._style_cache.get(cache_key)
        if cached is not None:
            return cached
        tag = (element.name or "").lower()
        classes = tuple(element.get("class", []) or [])
        class_set = frozenset(classes)
        element_id = element.get("id")
        candidates: List[StyleRule] = []
        seen_rules = set()
        candidate_groups: List[Iterable[StyleRule]] = [self.universal_rules]
        if tag:
            candidate_groups.append(self.rules_by_tag.get(tag, ()))
        if element_id:
            candidate_groups.append(self.rules_by_id.get(element_id, ()))
        for cls in classes:
            candidate_groups.append(self.rules_by_class.get(cls, ()))
        for group in candidate_groups:
            for rule in group:
                if rule.order in seen_rules:
                    continue
                seen_rules.add(rule.order)
                candidates.append(rule)
        candidates.sort(key=lambda rule: rule.order)
        style: Dict[str, str] = {}
        for rule in candidates:
            if rule.matches(tag, class_set, element_id):
                style.update(rule.declarations)
        inline = element.get("style")
        if inline:
            style.update(parse_declarations(inline))
        self._style_cache[cache_key] = style
        return style


def build_text_style(tag: str, style: Dict[str, str]) -> Dict[str, Any]:
    result: Dict[str, Any] = {}
    result["font_size"] = parse_font_size(style.get("font-size"), tag)
    result["font_family"] = style.get("font-family")
    weight = style.get("font-weight", "").lower()
    if weight in {"bold", "bolder"}:
        result["bold"] = True
    elif weight.isdigit():
        result["bold"] = int(weight) >= 600
    italic = style.get("font-style", "").lower()
    if italic:
        result["italic"] = italic == "italic"
    color = style.get("color")
    if color:
        result["color"] = color
    align = style.get("text-align", "").lower()
    if align in {"left", "center", "right", "justify"}:
        result["align"] = align
    line_height = parse_line_height(style.get("line-height"), result["font_size"])
    if line_height:
        result["line_height"] = line_height
    if style.get("letter-spacing"):
        result["letter_spacing"] = style.get("letter-spacing")
    if style.get("text-transform"):
        result["text_transform"] = style.get("text-transform")
    if style.get("white-space"):
        result["white_space"] = style.get("white-space")
    for side in ("left", "right", "top", "bottom"):
        padding = parse_length(style.get(f"padding-{side}"))
        if padding is not None:
            result[f"padding_{side}"] = padding
    return result


def css_weight_is_bold(weight: Optional[str]) -> bool:
    if weight is None:
        return False
    s = str(weight).strip().lower()
    if not s:
        return False
    if s.isdigit():
        try:
            return int(s) >= 600
        except ValueError:
            return False
    return s in {"bold", "bolder", "600", "700", "800", "900"}


def apply_text_style(base: Dict[str, Any], style: Dict[str, str], tag: Optional[str] = None) -> Dict[str, Any]:
    new_style = dict(base)
    size = parse_length(style.get("font-size"))
    if size:
        new_style["font_size"] = size
    if style.get("font-family"):
        new_style["font_family"] = style.get("font-family")
    weight = style.get("font-weight", "").lower()
    if weight in {"bold", "bolder"}:
        new_style["bold"] = True
    elif weight.isdigit():
        new_style["bold"] = int(weight) >= 600
    italic = style.get("font-style", "").lower()
    if italic:
        new_style["italic"] = italic == "italic"
    color = style.get("color")
    if color:
        new_style["color"] = color
    align = style.get("text-align", "").lower()
    if align in {"left", "center", "right", "justify"}:
        new_style["align"] = align
    line_height = parse_line_height(style.get("line-height"), new_style.get("font_size"))
    if line_height:
        new_style["line_height"] = line_height
    if style.get("letter-spacing"):
        new_style["letter_spacing"] = style.get("letter-spacing")
    if style.get("text-transform"):
        new_style["text_transform"] = style.get("text-transform")
    if style.get("white-space"):
        new_style["white_space"] = style.get("white-space")
    for side in ("left", "right", "top", "bottom"):
        padding = parse_length(style.get(f"padding-{side}"))
        if padding is not None:
            new_style[f"padding_{side}"] = padding
    if tag in {"strong", "b"}:
        new_style["bold"] = True
    if tag in {"em", "i"}:
        new_style["italic"] = True
    return new_style


def extract_layout(style: Dict[str, str], element: Optional[Tag] = None) -> LayoutBox:
    box = LayoutBox()
    attr_width = element.get("width") if element and element.has_attr("width") else None
    attr_height = element.get("height") if element and element.has_attr("height") else None
    box.width = parse_length(attr_width, SLIDE_REF_WIDTH)
    box.height = parse_length(attr_height, SLIDE_REF_HEIGHT)
    width_style = style.get("width")
    height_style = style.get("height")
    if width_style:
        box.width = parse_length(width_style, SLIDE_REF_WIDTH)
    if height_style:
        box.height = parse_length(height_style, SLIDE_REF_HEIGHT)
    box.left = parse_length(style.get("left"), SLIDE_REF_WIDTH)
    box.top = parse_length(style.get("top"), SLIDE_REF_HEIGHT)
    return box


def merge_runs(runs: List[TextRun]) -> List[TextRun]:
    merged: List[TextRun] = []
    for run in runs:
        if not run.text:
            continue
        if merged:
            prev = merged[-1]
            if (
                prev.font_size == run.font_size
                and prev.bold == run.bold
                and prev.italic == run.italic
                and prev.color == run.color
                and prev.font_family == run.font_family
            ):
                prev.text += run.text
                continue
        merged.append(run)
    return merged


def extract_text_runs(element: Tag, resolver: StyleResolver, base_style: Dict[str, Any]) -> List[TextRun]:
    runs: List[TextRun] = []

    def walk(node: Any, current_style: Dict[str, Any]) -> None:
        if isinstance(node, NavigableString):
            text = str(node)
            normalized = WHITESPACE_RE.sub(" ", text)
            if normalized.strip():
                transformed = apply_text_transform(normalized, current_style.get("text_transform"))
                runs.append(
                    TextRun(
                        text=transformed,
                        font_size=current_style.get("font_size"),
                        bold=current_style.get("bold"),
                        italic=current_style.get("italic"),
                        color=current_style.get("color"),
                        font_family=current_style.get("font_family"),
                    )
                )
            return
        if not isinstance(node, Tag):
            return
        if node.name == "br":
            runs.append(
                TextRun(
                    text="\n",
                    font_size=current_style.get("font_size"),
                    bold=current_style.get("bold"),
                    italic=current_style.get("italic"),
                    color=current_style.get("color"),
                )
            )
            return
        node_style = resolver.get_style(node)
        next_style = apply_text_style(current_style, node_style, tag=node.name)
        for child in node.children:
            walk(child, next_style)

    for child in element.children:
        walk(child, dict(base_style))
    if not runs:
        text = normalize_whitespace(element.get_text(" ", strip=True))
        if text:
            transformed = apply_text_transform(text, base_style.get("text_transform"))
            runs.append(
                TextRun(
                    text=transformed,
                    font_size=base_style.get("font_size"),
                    bold=base_style.get("bold"),
                    italic=base_style.get("italic"),
                    color=base_style.get("color"),
                    font_family=base_style.get("font_family"),
                )
            )
    return merge_runs(runs)


def detect_shape_style(style: Dict[str, str]) -> Dict[str, Any]:
    fill_color = style.get("background-color") or style.get("background")
    border_value = style.get("border", "")
    border_width = parse_length(style.get("border-width"))
    border_color = style.get("border-color")
    if border_value:
        for token in border_value.split():
            length = parse_length(token)
            if length is not None:
                border_width = length
            elif token.startswith("#") or token.startswith("rgb"):
                border_color = token
    border_radius = parse_length(style.get("border-radius"))
    side_info: Dict[str, Any] = {}
    for side in ("top", "right", "bottom", "left"):
        side_width = parse_length(style.get(f"border-{side}-width"))
        side_color = style.get(f"border-{side}-color")
        if side_width is not None:
            side_info[f"border_{side}_width"] = side_width
        if side_color:
            side_info[f"border_{side}_color"] = side_color
    if not any([fill_color, border_color, border_width, border_radius, side_info]):
        return {}
    result = {
        "fill_color": fill_color,
        "border_color": border_color,
        "border_width": border_width,
        "border_radius": border_radius,
    }
    result.update(side_info)
    return result


def parse_z_index(style: Dict[str, str]) -> int:
    z_str = (style.get("z-index") or "").strip().lower()
    if not z_str or z_str == "auto":
        return 0
    try:
        return int(float(z_str))
    except ValueError:
        return 0


def has_block_children(element: Tag) -> bool:
    block_tags = {"p", "div", "section", "article", "ul", "ol", "table", "h1", "h2", "h3", "h4", "figure"}
    for child in element.children:
        if isinstance(child, Tag) and (child.name or "").lower() in block_tags:
            return True
    return False


def normalize_image_source(value: str) -> str:
    raw = (value or "").strip()
    if not raw:
        return ""
    parsed = urlparse(raw)
    if parsed.scheme == "file":
        path_value = url2pathname(unquote(parsed.path or ""))
        if parsed.netloc:
            path_value = f"/{parsed.netloc}{path_value}"
        return path_value
    if parsed.scheme in {"http", "https"}:
        return unquote(parsed.path or raw)
    return unquote(raw)


def image_map_lookup_keys(src: str) -> List[str]:
    raw = (src or "").strip()
    if not raw:
        return []
    normalized = normalize_image_source(raw)
    candidates: List[str] = []
    for candidate in (
        raw,
        normalized,
        Path(normalized).name if normalized else "",
        Path(raw).name if raw else "",
    ):
        candidate = candidate.strip()
        if candidate:
            candidates.append(candidate)
            candidates.append(candidate.casefold())
    unique: List[str] = []
    seen = set()
    for candidate in candidates:
        if candidate in seen:
            continue
        seen.add(candidate)
        unique.append(candidate)
    return unique


def build_image_map(entries: Optional[List[Tuple[str, str]]]) -> ImageMap:
    mapping: ImageMap = {}
    for key, mapped_path in entries or []:
        for alias in image_map_lookup_keys(key):
            mapping[alias] = mapped_path
    return mapping


def resolve_image_mapping(src: str, image_map: Optional[ImageMap]) -> Optional[str]:
    if not src or not image_map:
        return None
    for key in image_map_lookup_keys(src):
        mapped = image_map.get(key)
        if mapped:
            return mapped
    return None


def resolve_image_path(src: str, base_dir: Path, image_map: Optional[ImageMap] = None) -> Optional[Path]:
    if not src:
        return None
    raw = resolve_image_mapping(src, image_map) or src.strip()
    parsed = urlparse(raw)
    cleaned = raw
    if parsed.scheme:
        scheme = parsed.scheme.lower()
        if scheme in {"http", "https", "data"}:
            return None
        if scheme == "file":
            path_value = url2pathname(unquote(parsed.path or ""))
            if parsed.netloc:
                path_value = f"/{parsed.netloc}{path_value}"
            cleaned = path_value
    else:
        cleaned = unquote(raw)

    path_candidates: List[Path] = []
    candidate_path = Path(cleaned).expanduser()
    if candidate_path.is_absolute():
        path_candidates.append(candidate_path)
    else:
        path_candidates.extend(
            [
                base_dir / candidate_path,
                Path.cwd() / candidate_path,
                base_dir / "assets" / candidate_path,
                base_dir / "images" / candidate_path,
                Path.cwd() / "assets" / candidate_path,
                Path.cwd() / "images" / candidate_path,
            ]
        )

    file_name = candidate_path.name or Path(cleaned).name
    if file_name:
        downloads_dir = Path.home() / "Downloads"
        path_candidates.append(downloads_dir / file_name)
        if "." not in file_name:
            for parent in (base_dir, base_dir / "assets", base_dir / "images", downloads_dir):
                if not parent.exists() or not parent.is_dir():
                    continue
                matches = sorted(parent.glob(f"{file_name}.*"))
                path_candidates.extend(matches[:3])

    seen: set[str] = set()
    for candidate in path_candidates:
        try:
            resolved = candidate.expanduser().resolve()
        except Exception:
            continue
        key = str(resolved)
        if key in seen:
            continue
        seen.add(key)
        if resolved.exists() and resolved.is_file():
            return resolved
    return None


def format_image_label(src: str) -> str:
    if not src:
        return "[画像]"
    parsed = urlparse(src)
    if parsed.scheme == "data":
        return "[埋め込み画像]"
    name = ""
    if parsed.scheme == "file":
        path_value = url2pathname(unquote(parsed.path or ""))
        name = Path(path_value).name
    elif parsed.scheme in {"http", "https"}:
        name = Path(unquote(parsed.path or "")).name
    else:
        name = Path(unquote(src)).name
    name = name or src.strip()
    return f"[画像] {name}" if name else "[画像]"


@lru_cache(maxsize=128)
def decode_data_url_image(data_url: str) -> Optional[bytes]:
    if not data_url.startswith("data:"):
        return None
    header, sep, payload = data_url.partition(",")
    if not sep:
        return None
    try:
        if ";base64" in header:
            return base64.b64decode(payload)
        return unquote(payload).encode("utf-8")
    except Exception:
        return None


def element_text(element: Tag) -> str:
    return element.get_text(" ", strip=True)


def extract_blocks(
    slide_el: Tag,
    resolver: StyleResolver,
    base_dir: Path,
    constraints_map: Dict[int, LayoutConstraint],
    image_map: Optional[ImageMap] = None,
) -> List[Block]:
    blocks: List[Block] = []
    order_counter = 0
    for element in slide_el.descendants:
        if not isinstance(element, Tag):
            continue
        if element.name in {"script", "style", "noscript"}:
            continue
        tag = element.name.lower()
        style = resolver.get_style(element)
        layout = extract_layout(style, element)
        text_style = build_text_style(tag, style)
        block: Optional[Block] = None

        if tag in {"h1", "h2", "h3", "h4", "p", "blockquote"}:
            runs = extract_text_runs(element, resolver, text_style)
            text = element_text(element)
            block = Block(kind="text", text=text, runs=runs, text_style=text_style, layout=layout)
            shape_style = detect_shape_style(style)
            if shape_style:
                block.shape_style = shape_style

        elif tag in {"ul", "ol"}:
            items = [li.get_text(" ", strip=True) for li in element.find_all("li", recursive=False)]
            items = [item for item in items if item]
            if items:
                block = Block(
                    kind="list",
                    items=items,
                    numbered=(tag == "ol"),
                    text_style=text_style,
                    layout=layout,
                )

        elif tag == "table":
            rows: List[List[str]] = []
            cell_styles: List[List[TableCell]] = []
            for tr in element.find_all("tr", recursive=False):
                row = []
                row_cells: List[TableCell] = []
                for cell in tr.find_all(["th", "td"], recursive=False):
                    cell_text = cell.get_text(" ", strip=True)
                    row.append(cell_text)
                    cell_style = resolver.get_style(cell)
                    cell_text_style = build_text_style(cell.name, cell_style)
                    row_cells.append(
                        TableCell(
                            text=cell_text,
                            background_color=cell_style.get("background-color") or cell_style.get("background"),
                            border_color=cell_style.get("border-color"),
                            border_width=parse_length(cell_style.get("border-width")),
                            text_style=cell_text_style,
                            vertical_align=cell_style.get("vertical-align"),
                            runs=extract_text_runs(cell, resolver, cell_text_style),
                            is_header=cell.name == "th",
                        )
                    )
                if row:
                    rows.append(row)
                    cell_styles.append(row_cells)
            if rows:
                block = Block(kind="table", table=rows, table_cells=cell_styles, text_style=text_style, layout=layout)

        elif tag == "img":
            src = element.get("src")
            mapped_src = resolve_image_mapping(src or "", image_map) or src or ""
            img_path = resolve_image_path(src, base_dir, image_map=image_map) if src else None
            block = Block(
                kind="image",
                image_path=img_path,
                image_alt=element.get("alt", "") or format_image_label(mapped_src),
                layout=layout,
            )

        elif tag in {"div", "span", "section"}:
            text_content = element_text(element)
            shape_style = detect_shape_style(style)
            if text_content and not has_block_children(element):
                runs = extract_text_runs(element, resolver, text_style)
                block = Block(kind="text", text=text_content, runs=runs, text_style=text_style, layout=layout)
                if shape_style:
                    block.shape_style = shape_style
            elif shape_style:
                block = Block(kind="shape", shape_style=shape_style, layout=layout)

        elif tag == "hr":
            block = Block(
                kind="shape",
                shape_style={"fill_color": style.get("background-color") or "#999999"},
                layout=layout,
            )

        if block:
            block.z_index = parse_z_index(style)
            block.order = order_counter
            order_counter += 1
            blocks.append(block)
            register_layout_constraints(element, block, resolver, constraints_map)
    return blocks


def estimate_block_height(block: Block) -> float:
    if block.layout.height:
        return block.layout.height
    style = block.text_style or {}
    font_size = style.get("font_size", DEFAULT_FONT_SIZE)
    if block.kind == "text":
        lines = max(1, block.text.count("\n") + 1)
        return max(80, lines * font_size * 1.4 + 20)
    if block.kind == "list":
        lines = max(1, len(block.items))
        return max(80, lines * font_size * 1.3 + 20)
    if block.kind == "table":
        rows = max(1, len(block.table))
        return rows * (font_size * 1.8)
    if block.kind == "image":
        return 240
    if block.kind == "shape":
        return 120
    if block.kind == "bar-chart":
        return 260
    return DEFAULT_BLOCK_HEIGHT


def assign_fallback_layouts(slides: List[SlideModel]) -> None:
    for slide in slides:
        slide_width = slide.canvas_width or SLIDE_REF_WIDTH
        flow_y = DEFAULT_PADDING_Y
        for block in slide.blocks:
            if block.layout.left is None:
                block.layout.left = DEFAULT_PADDING_X
            if block.layout.width is None:
                block.layout.width = max(slide_width - 2 * DEFAULT_PADDING_X, 200.0)
            block_height = block.layout.height or estimate_block_height(block)
            if block.layout.top is None:
                block.layout.top = flow_y
            block.layout.height = block_height
            flow_y = max(flow_y, block.layout.top + block.layout.height + FLOW_GAP)


def apply_layout_constraints(slides: List[SlideModel]) -> None:
    for slide in slides:
        if not slide.layout_constraints:
            continue
        constraints = sorted(slide.layout_constraints, key=lambda c: c.depth)
        for constraint in constraints:
            slots = [slot for slot in constraint.slots if slot.blocks]
            if len(slots) <= 1:
                continue
            if constraint.kind == "grid":
                apply_grid_constraint(constraint, slots)
            elif constraint.kind == "flex":
                apply_flex_constraint(constraint, slots)


def apply_grid_constraint(constraint: LayoutConstraint, slots: List[LayoutSlot]) -> None:
    container_left, container_top, container_width = compute_container_box(constraint)
    row_gap, col_gap = parse_gap_values(constraint.parent_style)
    column_widths = compute_grid_column_widths(
        container_width,
        constraint.parent_style.get("grid-template-columns"),
        col_gap,
        len(slots),
    )
    if not column_widths:
        column_widths = [container_width]
    col_count = max(1, len(column_widths))
    column_offsets: List[float] = []
    accum = 0.0
    for width in column_widths:
        column_offsets.append(accum)
        accum += width + col_gap
    row_top = container_top
    row_max_height = 0.0
    for idx, slot in enumerate(slots):
        col_index = idx % col_count
        if idx > 0 and col_index == 0:
            row_top += row_max_height + row_gap
            row_max_height = 0.0
        width = column_widths[col_index]
        slot_left = container_left + column_offsets[col_index]
        min_left, min_top, slot_width, slot_height = compute_slot_metrics(slot)
        slot_height = slot_height or DEFAULT_BLOCK_HEIGHT
        left_offset = slot_left - min_left
        top_offset = row_top - min_top
        for block in slot.blocks:
            orig_left = block.layout.left if block.layout.left is not None else min_left
            orig_top = block.layout.top if block.layout.top is not None else min_top
            orig_width = block.layout.width if block.layout.width is not None else slot_width
            block.layout.left = orig_left + left_offset
            block.layout.top = orig_top + top_offset
            block.layout.width = min(orig_width, width)
        row_max_height = max(row_max_height, slot_height)


def apply_flex_constraint(constraint: LayoutConstraint, slots: List[LayoutSlot]) -> None:
    container_left, container_top, container_width = compute_container_box(constraint)
    row_gap, col_gap = parse_gap_values(constraint.parent_style)
    direction = constraint.parent_style.get("flex-direction", "row").lower()
    metrics = [compute_slot_metrics(slot) for slot in slots]
    if direction.startswith("column"):
        top = container_top
        for slot, metric in zip(slots, metrics):
            min_left, min_top, slot_width, slot_height = metric
            slot_height = slot_height or DEFAULT_BLOCK_HEIGHT
            left_offset = container_left - min_left
            top_offset = top - min_top
            for block in slot.blocks:
                orig_left = block.layout.left if block.layout.left is not None else min_left
                orig_top = block.layout.top if block.layout.top is not None else min_top
                orig_width = block.layout.width if block.layout.width is not None else slot_width
                block.layout.left = orig_left + left_offset
                block.layout.top = orig_top + top_offset
                block.layout.width = min(orig_width, container_width)
            top += slot_height + row_gap
    else:
        width_available = max(container_width - col_gap * (len(slots) - 1), 50.0)
        width_per = width_available / len(slots) if slots else width_available
        left = container_left
        for slot, metric in zip(slots, metrics):
            min_left, min_top, slot_width, slot_height = metric
            left_offset = left - min_left
            top_offset = container_top - min_top
            for block in slot.blocks:
                orig_left = block.layout.left if block.layout.left is not None else min_left
                orig_top = block.layout.top if block.layout.top is not None else min_top
                orig_width = block.layout.width if block.layout.width is not None else slot_width
                block.layout.left = orig_left + left_offset
                block.layout.top = orig_top + top_offset
                block.layout.width = min(orig_width, width_per)
            left += width_per + col_gap


def is_footer_text_block(block: Block, slide_width: float, slide_height: float) -> bool:
    if block.kind != "text":
        return False
    top = block.layout.top or 0.0
    left = block.layout.left or 0.0
    width = block.layout.width or 0.0
    height = block.layout.height or 0.0
    if height <= 0 or width <= 0:
        return False
    if top < slide_height * 0.88 or height > 28:
        return False
    right = left + width
    if left <= slide_width * 0.3 or right >= slide_width * 0.7:
        return True
    text = block.text.strip().lower()
    return text.startswith("page ") or "contact" in text or text.startswith("©")


def resolve_bottom_text_overlaps(slide: SlideModel) -> None:
    width = slide.canvas_width or SLIDE_REF_WIDTH
    height = slide.canvas_height or SLIDE_REF_HEIGHT
    footers = [block for block in slide.blocks if is_footer_text_block(block, width, height)]
    if not footers:
        return
    footer_top = min((block.layout.top or 0.0) for block in footers)
    safe_gap = 6.0
    for block in slide.blocks:
        if block.kind != "text" or block in footers:
            continue
        if (block.text_style or {}).get("align") != "center":
            continue
        block_top = block.layout.top or 0.0
        block_height = block.layout.height or estimate_block_height(block)
        block_width = block.layout.width or 0.0
        if block_top < height * 0.75 or block_width < width * 0.45:
            continue
        block_bottom = block_top + block_height
        if block_bottom <= footer_top - safe_gap:
            continue
        block.layout.top = max(DEFAULT_PADDING_Y, footer_top - block_height - safe_gap)


def fit_slide_content(slide: SlideModel) -> None:
    if not slide.blocks:
        return
    width = slide.canvas_width or SLIDE_REF_WIDTH
    height = slide.canvas_height or SLIDE_REF_HEIGHT
    min_left = min((block.layout.left for block in slide.blocks if block.layout.left is not None), default=0.0)
    min_top = min((block.layout.top for block in slide.blocks if block.layout.top is not None), default=0.0)
    default_width = max(width - 2 * DEFAULT_PADDING_X, 200.0)
    max_right = max(
        (
            (block.layout.left or 0.0) + (block.layout.width or default_width)
            for block in slide.blocks
        ),
        default=width,
    )
    max_bottom = max(
        (
            (block.layout.top or 0.0) + (block.layout.height or estimate_block_height(block))
            for block in slide.blocks
        ),
        default=height,
    )
    content_width = max_right - min_left
    content_height = max_bottom - min_top
    avail_width = max(width - 2 * DEFAULT_PADDING_X, 200.0)
    avail_height = max(height - 2 * DEFAULT_PADDING_Y, 200.0)
    scale = 1.0
    needs_fit = False
    if content_width > avail_width + 1:
        scale = min(scale, avail_width / content_width)
        needs_fit = True
    if content_height > avail_height + 1:
        scale = min(scale, avail_height / content_height)
        needs_fit = True
    if min_left < DEFAULT_PADDING_X - 5 or min_top < DEFAULT_PADDING_Y - 5:
        needs_fit = True
    if max_right > width - DEFAULT_PADDING_X + 5 or max_bottom > height - DEFAULT_PADDING_Y + 5:
        needs_fit = True
    if not needs_fit:
        resolve_bottom_text_overlaps(slide)
        return
    if scale <= 0:
        scale = 1.0
    for block in slide.blocks:
        left = block.layout.left or 0.0
        top = block.layout.top or 0.0
        width_px = block.layout.width or (content_width if content_width > 0 else avail_width)
        height_px = block.layout.height or (content_height if content_height > 0 else avail_height)
        block.layout.left = DEFAULT_PADDING_X + (left - min_left) * scale
        block.layout.top = DEFAULT_PADDING_Y + (top - min_top) * scale
        block.layout.width = width_px * scale
        block.layout.height = height_px * scale
    resolve_bottom_text_overlaps(slide)


def parse_html_static(input_html: str, image_map: Optional[ImageMap] = None) -> List[SlideModel]:
    path = Path(input_html)
    base_dir = path.parent
    with path.open("r", encoding="utf-8") as f:
        soup = BeautifulSoup(f, "html.parser")
    resolver = StyleResolver(soup)
    slide_candidates = soup.select(".slide")
    if not slide_candidates:
        slide_candidates = soup.select("[data-slide]")
    if not slide_candidates:
        if soup.body:
            slide_candidates = [soup.body]
        else:
            raise ValueError("HTML に <body> が見つかりません。スライドを判定できません。")
    slides: List[SlideModel] = []
    for idx, slide_el in enumerate(slide_candidates):
        style = resolver.get_style(slide_el)
        background = style.get("background-color") or style.get("background")
        title_el = slide_el.find(["h1", "h2", "h3", "h4"])
        title = title_el.get_text(" ", strip=True) if title_el else None
        constraints_map: Dict[int, LayoutConstraint] = {}
        blocks = extract_blocks(slide_el, resolver, base_dir, constraints_map, image_map=image_map)
        constraints = sorted(constraints_map.values(), key=lambda c: c.depth)
        canvas_width = (
            parse_length(slide_el.get("width")) or parse_length(style.get("width")) or SLIDE_REF_WIDTH
        )
        canvas_height = (
            parse_length(slide_el.get("height"))
            or parse_length(style.get("height"))
            or parse_length(style.get("min-height"))
            or SLIDE_REF_HEIGHT
        )
        slides.append(
            SlideModel(
                title=title,
                background_color=background,
                blocks=blocks,
                layout_constraints=constraints,
                canvas_width=canvas_width,
                canvas_height=canvas_height,
            )
        )
    assign_fallback_layouts(slides)
    apply_layout_constraints(slides)
    for slide in slides:
        fit_slide_content(slide)
    return slides


def browser_block_to_model(
    block_info: Dict[str, Any],
    base_dir: Path,
    image_map: Optional[ImageMap] = None,
) -> Optional[Block]:
    kind = block_info.get("kind")
    rect = block_info.get("rect") or {}
    layout = LayoutBox(
        left=float(rect.get("x") or 0.0),
        top=float(rect.get("y") or 0.0),
        width=float(rect.get("w") or 0.0),
        height=float(rect.get("h") or 0.0),
    )
    block = Block(kind=kind or "text", layout=layout)
    block.z_index = int(block_info.get("zIndex", 0) or 0)
    block.order = int(block_info.get("order", 0) or 0)
    styles = block_info.get("styles") or {}
    align = (styles.get("textAlign") or "").lower() if styles.get("textAlign") else None
    if kind == "text":
        text = block_info.get("text", "")
        if not text.strip():
            return None
        base_font_size = float(styles.get("fontSizePx") or DEFAULT_FONT_SIZE)
        base_bold = css_weight_is_bold(styles.get("fontWeight"))
        base_italic = (styles.get("fontStyle") or "").lower() == "italic"
        base_color = styles.get("color")
        block.text = text
        block.text_style = {
            "font_size": base_font_size,
            "line_height": float(styles.get("lineHeightPx") or 0) or None,
            "bold": base_bold,
            "italic": base_italic,
            "color": base_color,
            "align": align,
            "font_family": styles.get("fontFamily"),
            "text_transform": styles.get("textTransform"),
            "white_space": styles.get("whiteSpace"),
            "padding_left": float(styles.get("paddingLeft") or 0),
            "padding_right": float(styles.get("paddingRight") or 0),
            "padding_top": float(styles.get("paddingTop") or 0),
            "padding_bottom": float(styles.get("paddingBottom") or 0),
        }
        run_entries = block_info.get("runs") or []
        runs: List[TextRun] = []
        for entry in run_entries:
            segment = entry.get("text", "")
            if not segment:
                continue
            run_font_size = float(entry.get("fontSizePx") or base_font_size)
            run_bold = css_weight_is_bold(entry.get("fontWeight")) if entry.get("fontWeight") else base_bold
            run_italic = (
                (entry.get("fontStyle") or "").lower() == "italic" if entry.get("fontStyle") else base_italic
            )
            run_color = entry.get("color") or base_color
            runs.append(
                TextRun(
                    text=segment,
                    font_size=run_font_size,
                    bold=run_bold,
                    italic=run_italic,
                    color=run_color,
                    font_family=entry.get("fontFamily") or styles.get("fontFamily"),
                )
            )
        if runs:
            block.runs = runs
        else:
            block.runs = [
                TextRun(
                    text=text,
                    font_size=base_font_size,
                    bold=base_bold,
                    italic=base_italic,
                    color=base_color,
                    font_family=styles.get("fontFamily"),
                )
            ]
    elif kind == "list":
        items = block_info.get("items") or []
        if not items:
            return None
        font_size = float(styles.get("fontSizePx") or DEFAULT_FONT_SIZE)
        bold = False
        fw = styles.get("fontWeight")
        if fw:
            fw_str = str(fw).lower()
            if fw_str.isdigit():
                bold = int(fw_str) >= 600
            elif fw_str in {"bold", "bolder"}:
                bold = True
        color = styles.get("color")
        block.items = items
        block.numbered = bool(block_info.get("ordered"))
        block.text_style = {
            "font_size": font_size,
            "line_height": float(styles.get("lineHeightPx") or 0) or None,
            "bold": bold,
            "italic": (styles.get("fontStyle") or "").lower() == "italic",
            "color": color,
            "align": align,
            "font_family": styles.get("fontFamily"),
            "text_transform": styles.get("textTransform"),
            "white_space": styles.get("whiteSpace"),
            "padding_left": float(styles.get("paddingLeft") or 0),
            "padding_right": float(styles.get("paddingRight") or 0),
            "padding_top": float(styles.get("paddingTop") or 0),
            "padding_bottom": float(styles.get("paddingBottom") or 0),
        }
        item_meta_entries = block_info.get("itemMeta") or []
        if item_meta_entries:
            items_meta: List[Dict[str, Any]] = []
            for item_entry in item_meta_entries:
                item_styles = item_entry.get("styles") or {}
                item_runs: List[TextRun] = []
                for run_info in item_entry.get("runs") or []:
                    run_text = run_info.get("text", "")
                    if not run_text:
                        continue
                    item_runs.append(
                        TextRun(
                            text=run_text,
                            font_size=float(
                                run_info.get("fontSizePx")
                                or item_styles.get("fontSizePx")
                                or block.text_style["font_size"]
                            ),
                            bold=css_weight_is_bold(run_info.get("fontWeight"))
                            if run_info.get("fontWeight") is not None
                            else css_weight_is_bold(item_styles.get("fontWeight")),
                            italic=(run_info.get("fontStyle") or "").lower() == "italic"
                            if run_info.get("fontStyle") is not None
                            else ((item_styles.get("fontStyle") or "").lower() == "italic"),
                            color=run_info.get("color") or item_styles.get("color") or color,
                            font_family=run_info.get("fontFamily")
                            or item_styles.get("fontFamily")
                            or styles.get("fontFamily"),
                        )
                    )
                items_meta.append(
                    {
                        "text": item_entry.get("text", ""),
                        "rect": item_entry.get("rect") or {},
                        "style": {
                            "font_size": float(item_styles.get("fontSizePx") or font_size),
                            "line_height": float(item_styles.get("lineHeightPx") or 0) or None,
                        "bold": css_weight_is_bold(item_styles.get("fontWeight")),
                        "italic": (item_styles.get("fontStyle") or "").lower() == "italic",
                        "color": item_styles.get("color") or color,
                        "align": (item_styles.get("textAlign") or "").lower() or align,
                        "font_family": item_styles.get("fontFamily") or styles.get("fontFamily"),
                        "text_transform": item_styles.get("textTransform"),
                        "white_space": item_styles.get("whiteSpace"),
                        "padding_left": float(item_styles.get("paddingLeft") or 0),
                        "padding_right": float(item_styles.get("paddingRight") or 0),
                        "padding_top": float(item_styles.get("paddingTop") or 0),
                        "padding_bottom": float(item_styles.get("paddingBottom") or 0),
                    },
                        "runs": item_runs,
                    }
                )
            block.vector_data = {
                "items_meta": items_meta,
                "padding_left": float(styles.get("paddingLeft") or 0),
                "padding_right": float(styles.get("paddingRight") or 0),
                "padding_top": float(styles.get("paddingTop") or 0),
                "padding_bottom": float(styles.get("paddingBottom") or 0),
                "list_style_type": styles.get("listStyleType"),
                "list_style_position": styles.get("listStylePosition"),
            }
    elif kind == "table":
        rows = block_info.get("rows") or []
        if not rows:
            return None
        block.table = rows
        block.text_style = {
            "font_size": float(styles.get("fontSizePx") or DEFAULT_FONT_SIZE * 0.7),
            "line_height": float(styles.get("lineHeightPx") or 0) or None,
            "color": styles.get("color"),
            "font_family": styles.get("fontFamily"),
        }
        row_cells_info = block_info.get("rowCells") or []
        table_cells: List[List[TableCell]] = []
        for row in row_cells_info:
            cell_objs: List[TableCell] = []
            for cell_info in row:
                cell_styles = cell_info.get("styles") or {}
                cell_runs: List[TextRun] = []
                for run_info in cell_info.get("runs") or []:
                    run_text = run_info.get("text", "")
                    if not run_text:
                        continue
                    cell_runs.append(
                        TextRun(
                            text=run_text,
                            font_size=float(
                                run_info.get("fontSizePx")
                                or cell_styles.get("fontSizePx")
                                or block.text_style["font_size"]
                            ),
                            bold=css_weight_is_bold(run_info.get("fontWeight"))
                            if run_info.get("fontWeight") is not None
                            else css_weight_is_bold(cell_styles.get("fontWeight")),
                            italic=(run_info.get("fontStyle") or "").lower() == "italic"
                            if run_info.get("fontStyle") is not None
                            else None,
                            color=run_info.get("color") or cell_styles.get("color"),
                            font_family=run_info.get("fontFamily") or cell_styles.get("fontFamily"),
                        )
                    )
                text_style = {
                    "font_size": float(cell_styles.get("fontSizePx") or block.text_style["font_size"]),
                    "line_height": float(cell_styles.get("lineHeightPx") or 0) or None,
                    "bold": css_weight_is_bold(cell_styles.get("fontWeight")),
                    "italic": (cell_styles.get("fontStyle") or "").lower() == "italic",
                    "color": cell_styles.get("color"),
                    "align": (cell_styles.get("textAlign") or "").lower() or None,
                    "font_family": cell_styles.get("fontFamily"),
                }
                cell_objs.append(
                    TableCell(
                        text=cell_info.get("text", ""),
                        background_color=cell_styles.get("backgroundColor"),
                        border_color=cell_styles.get("borderColor"),
                        border_width=float(cell_styles.get("borderWidth") or 0),
                        text_style=text_style,
                        vertical_align=(cell_styles.get("verticalAlign") or "").lower() or None,
                        runs=cell_runs,
                        is_header=bool(cell_info.get("isHeader")),
                        padding_left=float(cell_styles.get("paddingLeft") or 0),
                        padding_right=float(cell_styles.get("paddingRight") or 0),
                        padding_top=float(cell_styles.get("paddingTop") or 0),
                        padding_bottom=float(cell_styles.get("paddingBottom") or 0),
                    )
                )
            table_cells.append(cell_objs)
        if table_cells:
            block.table_cells = table_cells
        block.vector_data = {
            "column_widths": block_info.get("columnWidths") or [],
            "row_heights": block_info.get("rowHeights") or [],
        }
    elif kind == "image":
        src = block_info.get("src")
        if src:
            mapped_src = resolve_image_mapping(src, image_map) or src
            block.image_path = resolve_image_path(src, base_dir, image_map=image_map)
            block.image_alt = block_info.get("alt") or format_image_label(mapped_src)
        data_url = block_info.get("dataUrl")
        if data_url:
            block.vector_data["data_url"] = data_url
        if block_info.get("naturalWidth") and block_info.get("naturalHeight"):
            block.vector_data["natural_size"] = (
                float(block_info["naturalWidth"]),
                float(block_info["naturalHeight"]),
            )
    elif kind == "shape":
        shape = block_info.get("shape") or {}
        block.shape_style = {
            "fill_color": shape.get("fill"),
            "border_color": shape.get("borderColor"),
            "border_width": shape.get("borderWidth"),
            "border_radius": shape.get("borderRadius"),
            "is_ellipse": bool(shape.get("isEllipse")),
            "border_top_width": shape.get("borderTopWidth"),
            "border_right_width": shape.get("borderRightWidth"),
            "border_bottom_width": shape.get("borderBottomWidth"),
            "border_left_width": shape.get("borderLeftWidth"),
            "border_top_color": shape.get("borderTopColor"),
            "border_right_color": shape.get("borderRightColor"),
            "border_bottom_color": shape.get("borderBottomColor"),
            "border_left_color": shape.get("borderLeftColor"),
        }
    elif kind == "bar-chart":
        chart = block_info.get("chart") or {}
        rows = chart.get("rows") or []
        if not rows:
            return None
        block.kind = "bar-chart"
        block.text_style = {
            "font_size": float((chart.get("labelStyles") or {}).get("fontSizePx") or DEFAULT_FONT_SIZE * 0.7),
            "line_height": float((chart.get("labelStyles") or {}).get("lineHeightPx") or 0) or None,
            "bold": css_weight_is_bold((chart.get("labelStyles") or {}).get("fontWeight")),
            "italic": ((chart.get("labelStyles") or {}).get("fontStyle") or "").lower() == "italic",
            "color": (chart.get("labelStyles") or {}).get("color"),
            "align": ((chart.get("labelStyles") or {}).get("textAlign") or "").lower() or None,
            "font_family": (chart.get("labelStyles") or {}).get("fontFamily"),
        }
        block.vector_data = chart
    elif kind == "conic-gradient":
        gradient = block_info.get("gradient") or {}
        segments = gradient.get("segments") or []
        if not segments:
            return None
        block.kind = "conic-gradient"
        block.vector_data = {
            "segments": segments,
            "cx": gradient.get("cx"),
            "cy": gradient.get("cy"),
            "radius": gradient.get("radius"),
        }
    elif kind == "svg-polyline":
        poly = block_info.get("polyline") or {}
        points = poly.get("points") or []
        if len(points) < 2:
            return None
        block.kind = "vector_polyline"
        block.vector_data = {
            "points": points,
            "stroke": poly.get("stroke"),
            "stroke_width": poly.get("strokeWidth"),
            "fill": poly.get("fill"),
            "closed": bool(poly.get("closed")),
        }
    elif kind == "svg-circle":
        circle = block_info.get("circle") or {}
        block.kind = "vector_circle"
        block.vector_data = circle
    elif kind == "svg-ellipse":
        ellipse = block_info.get("ellipse") or {}
        block.kind = "vector_ellipse"
        block.vector_data = ellipse
    else:
        block.text = block_info.get("text", "")
    return block


def parse_html_browser(
    input_html: str,
    selector: str,
    viewport_w: int,
    viewport_h: int,
    dpi_scale: int,
    image_map: Optional[ImageMap] = None,
    rasterize_slides: str = "auto",
) -> List[SlideModel]:
    playwright_factory = ensure_playwright()
    if playwright_factory is None:
        raise RuntimeError("playwright がインストールされていません。pip install playwright を実行してください。")
    uri = file_or_url_to_uri(input_html)
    slides_data: List[Dict[str, Any]] = []
    with playwright_factory() as pw:
        browser = pw.chromium.launch()
        ctx = browser.new_context(
            viewport={"width": viewport_w, "height": viewport_h},
            device_scale_factor=dpi_scale,
        )
        page = ctx.new_page()
        page.goto(uri, wait_until="load")
        page.wait_for_timeout(150)
        slides_data = page.evaluate(BROWSER_COLLECT_JS, {"selector": selector})
        matched_slide_indexes = page.locator(selector).evaluate_all(
            """
            (elements) =>
              elements
                .map((element, index) => ({
                  index,
                  visible: element.offsetWidth > 0 && element.offsetHeight > 0
                }))
                .filter((item) => item.visible)
                .map((item) => item.index)
            """
        )
        ctx.close()
        browser.close()
    slides: List[SlideModel] = []
    base_dir = Path(input_html).parent
    for info in slides_data:
        rect = info.get("rect") or {}
        viewport = info.get("viewport") or {}
        canvas_w = float(rect.get("w") or viewport.get("w") or SLIDE_REF_WIDTH)
        canvas_h = float(rect.get("h") or viewport.get("h") or SLIDE_REF_HEIGHT)
        slide = SlideModel(
            title=None,
            background_color=info.get("background"),
            blocks=[],
            layout_constraints=[],
            canvas_width=canvas_w,
            canvas_height=canvas_h,
        )
        for block_info in info.get("blocks", []):
            block = browser_block_to_model(block_info, base_dir, image_map=image_map)
            if block:
                slide.blocks.append(block)
        slides.append(slide)
    target_indices = resolve_rasterize_slide_targets(rasterize_slides, slides)
    if target_indices:
        with playwright_factory() as pw:
            browser = pw.chromium.launch()
            ctx = browser.new_context(
                viewport={"width": viewport_w, "height": viewport_h},
                device_scale_factor=dpi_scale,
            )
            page = ctx.new_page()
            page.goto(uri, wait_until="load")
            page.wait_for_timeout(150)
            if matched_slide_indexes:
                slide_locator = page.locator(selector)
                for slide_index, reason in sorted(target_indices.items()):
                    if slide_index < 0 or slide_index >= len(matched_slide_indexes) or slide_index >= len(slides):
                        continue
                    image_bytes = slide_locator.nth(matched_slide_indexes[slide_index]).screenshot(
                        type="png",
                        animations="disabled",
                    )
                    slides[slide_index].raster_image_data_url = encode_png_data_url(image_bytes)
                    slides[slide_index].raster_reason = reason
            elif slides and 0 in target_indices:
                image_bytes = page.locator("body").screenshot(type="png", animations="disabled")
                slides[0].raster_image_data_url = encode_png_data_url(image_bytes)
                slides[0].raster_reason = target_indices[0]
            ctx.close()
            browser.close()
    for slide in slides:
        resolve_bottom_text_overlaps(slide)
    return slides


def ensure_slide_transform(slide_model: SlideModel, prs: Presentation) -> None:
    if slide_model.scale_x is not None and slide_model.scale_y is not None:
        return
    canvas_width = slide_model.canvas_width or SLIDE_REF_WIDTH
    canvas_height = slide_model.canvas_height or SLIDE_REF_HEIGHT
    if canvas_width <= 0:
        canvas_width = SLIDE_REF_WIDTH
    if canvas_height <= 0:
        canvas_height = SLIDE_REF_HEIGHT
    scale_x = prs.slide_width / canvas_width
    scale_y = prs.slide_height / canvas_height
    if slide_model.transform_mode == "fill":
        slide_model.scale_x = scale_x or 1.0
        slide_model.scale_y = scale_y or 1.0
        slide_model.scale = min(slide_model.scale_x, slide_model.scale_y)
        slide_model.offset_x = 0.0
        slide_model.offset_y = 0.0
        return
    scale = min(scale_x, scale_y)
    if scale <= 0:
        scale = scale_x or scale_y or 1.0
    slide_model.scale = scale
    slide_model.scale_x = scale
    slide_model.scale_y = scale
    slide_model.offset_x = (prs.slide_width - canvas_width * scale) / 2
    slide_model.offset_y = (prs.slide_height - canvas_height * scale) / 2


def length_to_emu(value_px: float, axis: str, prs: Presentation, slide_model: SlideModel) -> int:
    ensure_slide_transform(slide_model, prs)
    scale = (slide_model.scale_x if axis == "x" else slide_model.scale_y) or slide_model.scale or 1.0
    return int(max(value_px, 0.0) * scale)


def position_to_emu(value_px: float, axis: str, prs: Presentation, slide_model: SlideModel) -> int:
    ensure_slide_transform(slide_model, prs)
    scale = (slide_model.scale_x if axis == "x" else slide_model.scale_y) or slide_model.scale or 1.0
    base = slide_model.offset_x if axis == "x" else slide_model.offset_y
    return int(base + value_px * scale)


def apply_paragraph_alignment(paragraph, align: Optional[str]) -> None:
    if not align:
        return
    align_map = {
        "left": PP_ALIGN.LEFT,
        "start": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
        "end": PP_ALIGN.RIGHT,
        "justify": PP_ALIGN.JUSTIFY,
    }
    paragraph.alignment = align_map.get(align, PP_ALIGN.LEFT)


def set_font_typeface(font, family: str) -> None:
    try:
        r_pr = font._element
    except Exception:
        return
    for tag in ("a:latin", "a:ea", "a:cs"):
        child = r_pr.find(qn(tag))
        if child is None:
            child = OxmlElement(tag)
            r_pr.append(child)
        child.set("typeface", family)


def apply_font_style(font, style: Dict[str, Any]) -> None:
    size = style.get("font_size")
    if size is not None:
        font.size = Pt(px_to_pt(size))
    if style.get("bold") is not None:
        font.bold = bool(style["bold"])
    if style.get("italic") is not None:
        font.italic = bool(style["italic"])
    color = style.get("color")
    rgb = css_color_to_rgb_tuple(color)
    if rgb:
        font.color.rgb = RGBColor(*rgb)
    family = style.get("font_family")
    if family:
        resolved_family = resolve_font_family_name(family) or family.split(",")[0].strip().strip("'\"")
        font.name = resolved_family
        set_font_typeface(font, resolved_family)


def reset_paragraph_spacing(paragraph) -> None:
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)


def apply_paragraph_style(paragraph, style: Dict[str, Any]) -> None:
    apply_paragraph_alignment(paragraph, style.get("align"))
    reset_paragraph_spacing(paragraph)
    line_height = style.get("line_height")
    if line_height:
        try:
            paragraph.line_spacing = Pt(px_to_pt(line_height))
        except Exception:
            pass


def set_text_frame_padding(text_frame, left: int = 0, right: int = 0, top: int = 0, bottom: int = 0) -> None:
    text_frame.margin_left = left
    text_frame.margin_right = right
    text_frame.margin_top = top
    text_frame.margin_bottom = bottom


def margin_from_padding_px(padding_px: Optional[float], default_pt: float) -> int:
    if padding_px is None or padding_px <= 0:
        return Pt(default_pt)
    return Pt(max(0.0, px_to_pt(padding_px)))


def apply_text_frame_box_model(text_frame, style: Dict[str, Any], default_pt: float = 0.0) -> None:
    set_text_frame_padding(
        text_frame,
        margin_from_padding_px(style.get("padding_left"), default_pt),
        margin_from_padding_px(style.get("padding_right"), default_pt),
        margin_from_padding_px(style.get("padding_top"), default_pt),
        margin_from_padding_px(style.get("padding_bottom"), default_pt),
    )
    white_space = (style.get("white_space") or "").strip().lower()
    text_frame.word_wrap = white_space not in {"nowrap", "pre"}


def populate_text_frame(text_frame, runs: List[TextRun], base_style: Dict[str, Any], default_text: str = "") -> None:
    existing_wrap = text_frame.word_wrap
    text_frame.clear()
    if existing_wrap is not None:
        text_frame.word_wrap = existing_wrap
    normalized_runs = runs or [
        TextRun(
            text=default_text,
            font_size=base_style.get("font_size"),
            bold=base_style.get("bold"),
            italic=base_style.get("italic"),
            color=base_style.get("color"),
            font_family=base_style.get("font_family"),
        )
    ]
    paragraph = text_frame.paragraphs[0]
    apply_paragraph_style(paragraph, base_style)
    for run in normalized_runs:
        pieces = run.text.split("\n")
        for idx, piece in enumerate(pieces):
            if idx > 0:
                paragraph = text_frame.add_paragraph()
                apply_paragraph_style(paragraph, base_style)
            ppt_run = paragraph.add_run()
            ppt_run.text = piece
            run_style = dict(base_style)
            if run.font_size is not None:
                run_style["font_size"] = run.font_size
            if run.bold is not None:
                run_style["bold"] = run.bold
            if run.italic is not None:
                run_style["italic"] = run.italic
            if run.color is not None:
                run_style["color"] = run.color
            if run.font_family is not None:
                run_style["font_family"] = run.font_family
            apply_font_style(ppt_run.font, run_style)


def clone_text_runs(runs: List[TextRun]) -> List[TextRun]:
    return [
        TextRun(
            text=run.text,
            font_size=run.font_size,
            bold=run.bold,
            italic=run.italic,
            color=run.color,
            font_family=run.font_family,
        )
        for run in runs
    ]


def prepend_prefix_to_runs(runs: List[TextRun], prefix: str, base_style: Dict[str, Any]) -> List[TextRun]:
    if not prefix:
        return clone_text_runs(runs)
    prepared = clone_text_runs(runs)
    if prepared:
        first = prepared[0]
        prepared[0] = TextRun(
            text=prefix + first.text,
            font_size=first.font_size,
            bold=first.bold,
            italic=first.italic,
            color=first.color,
            font_family=first.font_family,
        )
        return prepared
    return [
        TextRun(
            text=prefix,
            font_size=base_style.get("font_size"),
            bold=base_style.get("bold"),
            italic=base_style.get("italic"),
            color=base_style.get("color"),
            font_family=base_style.get("font_family"),
        )
    ]


def fit_relative_lengths(total: int, sizes: List[float], count: int) -> List[int]:
    if count <= 0:
        return []
    normalized = [max(float(size or 0), 0.0) for size in sizes[:count]]
    if len(normalized) < count:
        normalized.extend([0.0] * (count - len(normalized)))
    if not any(normalized):
        normalized = [1.0] * count
    overall = sum(normalized) or float(count)
    boundaries = [0]
    cumulative = 0.0
    for idx in range(count - 1):
        cumulative += normalized[idx]
        boundaries.append(int(round(total * cumulative / overall)))
    boundaries.append(total)
    lengths = [max(1, boundaries[idx + 1] - boundaries[idx]) for idx in range(count)]
    diff = total - sum(lengths)
    lengths[-1] += diff
    return lengths


def add_border_line(slide, x1: float, y1: float, x2: float, y2: float, color: str, width_px: float, prs, slide_model) -> None:
    rgb = css_color_to_rgb_tuple(color)
    if not rgb or width_px <= 0:
        return
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        position_to_emu(x1, "x", prs, slide_model),
        position_to_emu(y1, "y", prs, slide_model),
        position_to_emu(x2, "x", prs, slide_model),
        position_to_emu(y2, "y", prs, slide_model),
    )
    connector.line.color.rgb = RGBColor(*rgb)
    connector.line.width = Pt(px_to_pt(width_px))


def add_text_block(slide, block: Block, prs: Presentation, slide_model: SlideModel) -> None:
    fallback_width = max((slide_model.canvas_width or SLIDE_REF_WIDTH) - 2 * DEFAULT_PADDING_X, 200.0)
    left_px = block.layout.left or DEFAULT_PADDING_X
    top_px = block.layout.top or DEFAULT_PADDING_Y
    width_px = block.layout.width or fallback_width
    if (
        not block.shape_style
        and len(block.text.strip()) <= 64
        and (block.layout.height or 0) <= 90
        and width_px <= (slide_model.canvas_width or SLIDE_REF_WIDTH) * 0.7
    ):
        extra_ratio = 0.18 if len(block.text.strip()) <= 48 else 0.12
        extra_width = width_px * extra_ratio
        max_right = slide_model.canvas_width or SLIDE_REF_WIDTH
        if left_px + width_px + extra_width <= max_right:
            width_px += extra_width
    left = position_to_emu(left_px, "x", prs, slide_model)
    top = position_to_emu(top_px, "y", prs, slide_model)
    width = length_to_emu(width_px, "x", prs, slide_model)
    height = length_to_emu(block.layout.height or estimate_block_height(block), "y", prs, slide_model)
    box = slide.shapes.add_textbox(left, top, width, height)
    base_style = block.text_style or {}
    apply_text_frame_box_model(box.text_frame, base_style)
    populate_text_frame(box.text_frame, block.runs, base_style, default_text=block.text)
    if (
        (base_style.get("white_space") or "").strip().lower() not in {"nowrap", "pre"}
        and "\n" not in block.text
        and len(block.text.strip()) <= 32
        and (block.layout.height or 0) <= 70
    ):
        box.text_frame.word_wrap = False
    if block.shape_style:
        shape_style = block.shape_style
        fill_color = css_color_to_rgb_tuple(shape_style.get("fill_color"))
        if fill_color:
            box.fill.solid()
            box.fill.fore_color.rgb = RGBColor(*fill_color)
        else:
            box.fill.background()
        border_color = css_color_to_rgb_tuple(shape_style.get("border_color"))
        if border_color:
            line = box.line
            line.color.rgb = RGBColor(*border_color)
            width = shape_style.get("border_width") or 2
            line.width = Pt(px_to_pt(width))


def add_list_block(slide, block: Block, prs: Presentation, slide_model: SlideModel) -> None:
    base_style = block.text_style or {}
    item_meta = block.vector_data.get("items_meta") or []
    list_style_type = (block.vector_data.get("list_style_type") or "").lower()
    if item_meta:
        show_bullet = not block.numbered and list_style_type not in {"none", ""}
        for idx, item in enumerate(item_meta):
            rect = item.get("rect") or {}
            item_style = dict(base_style)
            item_style.update({k: v for k, v in (item.get("style") or {}).items() if v is not None})
            item_runs = item.get("runs") or []
            prefix = f"{idx + 1}. " if block.numbered else ("• " if show_bullet else "")
            prepared_runs = prepend_prefix_to_runs(item_runs, prefix, item_style)
            default_text = prefix + (item.get("text") or "").strip()
            left_px = (block.layout.left or DEFAULT_PADDING_X) + float(rect.get("x") or 0.0)
            top_px = (block.layout.top or DEFAULT_PADDING_Y) + float(rect.get("y") or 0.0)
            width_px = float(rect.get("w") or block.layout.width or 200.0)
            height_px = float(rect.get("h") or estimate_block_height(block))
            box = slide.shapes.add_textbox(
                position_to_emu(left_px, "x", prs, slide_model),
                position_to_emu(top_px, "y", prs, slide_model),
                length_to_emu(width_px, "x", prs, slide_model),
                length_to_emu(height_px, "y", prs, slide_model),
            )
            tf = box.text_frame
            apply_text_frame_box_model(tf, item_style)
            populate_text_frame(tf, prepared_runs, item_style, default_text=default_text)
            box.fill.background()
            box.line.fill.background()
        return
    left = position_to_emu(
        block.layout.left or DEFAULT_PADDING_X + DEFAULT_LIST_INDENT, "x", prs, slide_model
    )
    top = position_to_emu(block.layout.top or DEFAULT_PADDING_Y, "y", prs, slide_model)
    fallback_width = max(
        (slide_model.canvas_width or SLIDE_REF_WIDTH) - 2 * DEFAULT_PADDING_X - DEFAULT_LIST_INDENT,
        200.0,
    )
    width = length_to_emu(block.layout.width or fallback_width, "x", prs, slide_model)
    height = length_to_emu(block.layout.height or estimate_block_height(block), "y", prs, slide_model)
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    apply_text_frame_box_model(tf, base_style)
    tf.clear()
    for idx, item in enumerate(block.items):
        paragraph = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        text = item.strip()
        if block.numbered:
            text = f"{idx + 1}. {text}"
        else:
            text = f"• {text}"
        paragraph.text = text
        apply_paragraph_style(paragraph, base_style)
        apply_font_style(paragraph.font, {"font_size": base_style.get("font_size", DEFAULT_FONT_SIZE), **base_style})


def apply_cell_border(cell, color: Optional[str], width_px: float) -> None:
    if not color or css_is_transparent(color):
        return
    rgb = css_color_to_rgb_tuple(color)
    if not rgb:
        return
    width_pt = px_to_pt(width_px or 1.0)
    width_emu = Pt(width_pt).emu
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = tcPr.find(qn(tag))
        if ln is None:
            ln = OxmlElement(tag)
            tcPr.append(ln)
        ln.set("w", str(int(width_emu)))
        solid = ln.find(qn("a:solidFill"))
        if solid is None:
            solid = OxmlElement("a:solidFill")
            ln.append(solid)
        srgb = solid.find(qn("a:srgbClr"))
        if srgb is None:
            srgb = OxmlElement("a:srgbClr")
            solid.append(srgb)
        srgb.set("val", "{:02X}{:02X}{:02X}".format(*rgb))
        prst = ln.find(qn("a:prstDash"))
        if prst is None:
            prst = OxmlElement("a:prstDash")
            prst.set("val", "solid")
            ln.append(prst)


def add_table_block(slide, block: Block, prs: Presentation, slide_model: SlideModel) -> None:
    left = position_to_emu(block.layout.left or DEFAULT_PADDING_X, "x", prs, slide_model)
    top = position_to_emu(block.layout.top or DEFAULT_PADDING_Y, "y", prs, slide_model)
    fallback_width = max((slide_model.canvas_width or SLIDE_REF_WIDTH) - 2 * DEFAULT_PADDING_X, 200.0)
    width = length_to_emu(block.layout.width or fallback_width, "x", prs, slide_model)
    height = length_to_emu(block.layout.height or estimate_block_height(block), "y", prs, slide_model)
    rows = len(block.table)
    cols = max(len(r) for r in block.table)
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    column_widths = fit_relative_lengths(width, block.vector_data.get("column_widths") or [], cols)
    row_heights = fit_relative_lengths(height, block.vector_data.get("row_heights") or [], rows)
    for c in range(cols):
        table.columns[c].width = column_widths[c]
    for r in range(rows):
        table.rows[r].height = row_heights[r]
    base_style = block.text_style or {}
    for r, row in enumerate(block.table):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell_info = None
            if block.table_cells and r < len(block.table_cells):
                row_cells = block.table_cells[r]
                if c < len(row_cells):
                    cell_info = row_cells[c]
            text_value = value
            if cell_info and cell_info.text:
                text_value = cell_info.text
            tf = cell.text_frame
            style = dict(base_style)
            if cell_info:
                cell_style = cell_info.text_style or {}
                style.update({k: v for k, v in cell_style.items() if v is not None})
            if cell_info and cell_info.is_header and style.get("bold") is None:
                style["bold"] = True
            cell.margin_left = margin_from_padding_px(cell_info.padding_left if cell_info else None, 6)
            cell.margin_right = margin_from_padding_px(cell_info.padding_right if cell_info else None, 6)
            cell.margin_top = margin_from_padding_px(cell_info.padding_top if cell_info else None, 3)
            cell.margin_bottom = margin_from_padding_px(cell_info.padding_bottom if cell_info else None, 3)
            if cell_info and cell_info.vertical_align:
                vertical_align = cell_info.vertical_align
                if vertical_align in {"middle", "center"}:
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                elif vertical_align == "bottom":
                    cell.vertical_anchor = MSO_ANCHOR.BOTTOM
                else:
                    cell.vertical_anchor = MSO_ANCHOR.TOP
            populate_text_frame(tf, cell_info.runs if cell_info else [], style, default_text=text_value)
            if cell_info:
                bg = cell_info.background_color
                if bg and not css_is_transparent(bg):
                    rgb_bg = css_color_to_rgb_tuple(bg)
                    if rgb_bg:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor(*rgb_bg)
                else:
                    cell.fill.background()
                if cell_info.border_color:
                    apply_cell_border(cell, cell_info.border_color, cell_info.border_width or 1.0)


def add_rasterized_slide_block(slide, slide_model: SlideModel, prs: Presentation) -> bool:
    if not slide_model.raster_image_data_url:
        return False
    image_bytes = decode_data_url_image(slide_model.raster_image_data_url)
    if not image_bytes:
        return False
    ensure_slide_transform(slide_model, prs)
    left = position_to_emu(0, "x", prs, slide_model)
    top = position_to_emu(0, "y", prs, slide_model)
    width = length_to_emu(slide_model.canvas_width or SLIDE_REF_WIDTH, "x", prs, slide_model)
    height = length_to_emu(slide_model.canvas_height or SLIDE_REF_HEIGHT, "y", prs, slide_model)
    slide.shapes.add_picture(io.BytesIO(image_bytes), left, top, width=width, height=height)
    return True


@lru_cache(maxsize=256)
def _read_image_size(image_path: str) -> Optional[Tuple[int, int]]:
    image_module, _ = ensure_pillow()
    if not image_module:
        return None
    try:
        with image_module.open(image_path) as img:
            return img.size
    except Exception:
        return None


def _read_image_size_from_bytes(image_bytes: bytes) -> Optional[Tuple[int, int]]:
    image_module, _ = ensure_pillow()
    if not image_module:
        return None
    try:
        with image_module.open(io.BytesIO(image_bytes)) as img:
            return img.size
    except Exception:
        return None


def resolve_image_size(block: Block) -> Tuple[float, float]:
    width = block.layout.width or 640
    height = block.layout.height or 360
    if block.image_path:
        natural_size = _read_image_size(str(block.image_path))
        if natural_size:
            natural_w, natural_h = natural_size
            if block.layout.width and block.layout.height:
                return block.layout.width, block.layout.height
            if block.layout.width and not block.layout.height:
                ratio = block.layout.width / natural_w
                return block.layout.width, natural_h * ratio
            if block.layout.height and not block.layout.width:
                ratio = block.layout.height / natural_h
                return natural_w * ratio, block.layout.height
            return natural_w, natural_h
    if block.vector_data:
        natural_size = block.vector_data.get("natural_size")
        if natural_size:
            natural_w, natural_h = natural_size
            if block.layout.width and not block.layout.height:
                ratio = block.layout.width / natural_w
                return block.layout.width, natural_h * ratio
            if block.layout.height and not block.layout.width:
                ratio = block.layout.height / natural_h
                return natural_w * ratio, block.layout.height
        data_url = block.vector_data.get("data_url")
        if data_url:
            image_bytes = decode_data_url_image(data_url)
            if image_bytes:
                loaded_size = _read_image_size_from_bytes(image_bytes)
                if loaded_size:
                    natural_w, natural_h = loaded_size
                    if block.layout.width and not block.layout.height:
                        ratio = block.layout.width / natural_w
                        return block.layout.width, natural_h * ratio
                    if block.layout.height and not block.layout.width:
                        ratio = block.layout.height / natural_h
                        return natural_w * ratio, block.layout.height
                    if not block.layout.width and not block.layout.height:
                        return natural_w, natural_h
    return width, height


def add_image_block(slide, block: Block, prs: Presentation, slide_model: SlideModel) -> None:
    width_px, height_px = resolve_image_size(block)
    left = position_to_emu(block.layout.left or DEFAULT_PADDING_X, "x", prs, slide_model)
    top = position_to_emu(block.layout.top or DEFAULT_PADDING_Y, "y", prs, slide_model)
    width = length_to_emu(width_px, "x", prs, slide_model)
    height = length_to_emu(height_px, "y", prs, slide_model)
    image_bytes = decode_data_url_image(block.vector_data.get("data_url", "")) if block.vector_data else None
    if block.image_path and block.image_path.exists():
        slide.shapes.add_picture(str(block.image_path), left, top, width=width, height=height)
    elif image_bytes:
        slide.shapes.add_picture(io.BytesIO(image_bytes), left, top, width=width, height=height)
    else:
        placeholder = slide.shapes.add_textbox(left, top, width, height)
        tf = placeholder.text_frame
        tf.word_wrap = True
        tf.clear()
        title = tf.paragraphs[0]
        title.text = "画像未解決"
        title.alignment = PP_ALIGN.CENTER
        reset_paragraph_spacing(title)
        apply_font_style(title.font, {"font_size": 22, "bold": True, "color": "#666666"})
        caption = tf.add_paragraph()
        caption.text = block.image_alt or "[画像]"
        caption.alignment = PP_ALIGN.CENTER
        reset_paragraph_spacing(caption)
        apply_font_style(caption.font, {"font_size": 12, "color": "#888888"})
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        placeholder.fill.solid()
        placeholder.fill.fore_color.rgb = RGBColor(246, 246, 246)
        placeholder.line.color.rgb = RGBColor(200, 200, 200)
        placeholder.line.width = Pt(px_to_pt(2))


def add_shape_block(slide, block: Block, prs: Presentation, slide_model: SlideModel) -> None:
    left = position_to_emu(block.layout.left or DEFAULT_PADDING_X, "x", prs, slide_model)
    top = position_to_emu(block.layout.top or DEFAULT_PADDING_Y, "y", prs, slide_model)
    width = length_to_emu(block.layout.width or 240, "x", prs, slide_model)
    height = length_to_emu(block.layout.height or 120, "y", prs, slide_model)
    shape_type = (
        MSO_AUTO_SHAPE_TYPE.OVAL
        if block.shape_style.get("is_ellipse")
        else (
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
            if block.shape_style.get("border_radius")
            else MSO_AUTO_SHAPE_TYPE.RECTANGLE
        )
    )
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    fill_color = css_color_to_rgb_tuple(block.shape_style.get("fill_color"))
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*fill_color)
    else:
        shape.fill.background()
    border_entries = {}
    for side in ("top", "right", "bottom", "left"):
        width_px = float(block.shape_style.get(f"border_{side}_width") or 0.0)
        color = block.shape_style.get(f"border_{side}_color") or block.shape_style.get("border_color")
        if width_px > 0 and color and not css_is_transparent(color):
            border_entries[side] = (width_px, color)
    if border_entries:
        unique_borders = {(round(width_px, 3), color) for width_px, color in border_entries.values()}
        if len(border_entries) == 4 and len(unique_borders) == 1:
            width_px, color = next(iter(border_entries.values()))
            border_color = css_color_to_rgb_tuple(color)
            if border_color:
                shape.line.color.rgb = RGBColor(*border_color)
                shape.line.width = Pt(px_to_pt(width_px))
            else:
                shape.line.fill.background()
                shape.line.width = 0
        else:
            shape.line.fill.background()
            shape.line.width = 0
            x = block.layout.left or DEFAULT_PADDING_X
            y = block.layout.top or DEFAULT_PADDING_Y
            w = block.layout.width or 240
            h = block.layout.height or 120
            if "top" in border_entries:
                add_border_line(slide, x, y, x + w, y, border_entries["top"][1], border_entries["top"][0], prs, slide_model)
            if "right" in border_entries:
                add_border_line(slide, x + w, y, x + w, y + h, border_entries["right"][1], border_entries["right"][0], prs, slide_model)
            if "bottom" in border_entries:
                add_border_line(slide, x, y + h, x + w, y + h, border_entries["bottom"][1], border_entries["bottom"][0], prs, slide_model)
            if "left" in border_entries:
                add_border_line(slide, x, y, x, y + h, border_entries["left"][1], border_entries["left"][0], prs, slide_model)
    else:
        border_color = css_color_to_rgb_tuple(block.shape_style.get("border_color"))
        border_width = float(block.shape_style.get("border_width") or 0.0)
        if border_color and border_width > 0:
            shape.line.color.rgb = RGBColor(*border_color)
            width_px = border_width or 2
            shape.line.width = Pt(px_to_pt(width_px))
        else:
            shape.line.fill.background()
            shape.line.width = 0


def add_bar_chart_block(slide, block: Block, prs: Presentation, slide_model: SlideModel) -> None:
    chart = block.vector_data or {}
    rows = chart.get("rows") or []
    if not rows:
        return
    label_style = dict(block.text_style or {})
    title = chart.get("title") or {}
    title_text = (title.get("text") or "").strip()
    if title_text:
        title_rect = title.get("rect") or {}
        title_style = title.get("styles") or {}
        title_box = slide.shapes.add_textbox(
            position_to_emu((block.layout.left or 0) + float(title_rect.get("x") or 0), "x", prs, slide_model),
            position_to_emu((block.layout.top or 0) + float(title_rect.get("y") or 0), "y", prs, slide_model),
            length_to_emu(float(title_rect.get("w") or block.layout.width or 160.0), "x", prs, slide_model),
            length_to_emu(float(title_rect.get("h") or 18.0), "y", prs, slide_model),
        )
        set_text_frame_padding(title_box.text_frame)
        populate_text_frame(
            title_box.text_frame,
            [
                TextRun(
                    text=title_text,
                    font_size=float(title_style.get("fontSizePx") or label_style.get("font_size") or DEFAULT_FONT_SIZE * 0.55),
                    bold=css_weight_is_bold(title_style.get("fontWeight")),
                    italic=(title_style.get("fontStyle") or "").lower() == "italic",
                    color=title_style.get("color"),
                    font_family=title_style.get("fontFamily"),
                )
            ],
            {
                "font_size": float(title_style.get("fontSizePx") or label_style.get("font_size") or DEFAULT_FONT_SIZE * 0.55),
                "bold": css_weight_is_bold(title_style.get("fontWeight")),
                "italic": (title_style.get("fontStyle") or "").lower() == "italic",
                "color": title_style.get("color"),
                "align": (title_style.get("textAlign") or "left").lower(),
                "font_family": title_style.get("fontFamily"),
            },
            default_text=title_text,
        )
        title_box.fill.background()
        title_box.line.fill.background()
    for row in rows:
        label_box = slide.shapes.add_textbox(
            position_to_emu(block.layout.left or 0, "x", prs, slide_model),
            position_to_emu((block.layout.top or 0) + float((row.get("rowRect") or {}).get("y") or 0), "y", prs, slide_model),
            length_to_emu(float(chart.get("labelColumnWidth") or 180.0), "x", prs, slide_model),
            length_to_emu(float((row.get("rowRect") or {}).get("h") or 28.0), "y", prs, slide_model),
        )
        set_text_frame_padding(label_box.text_frame)
        populate_text_frame(
            label_box.text_frame,
            [
                TextRun(
                    text=row.get("label", ""),
                    font_size=label_style.get("font_size"),
                    bold=label_style.get("bold"),
                    italic=label_style.get("italic"),
                    color=label_style.get("color"),
                    font_family=label_style.get("font_family"),
                )
            ],
            {**label_style, "align": "left"},
            default_text=row.get("label", ""),
        )
        label_box.fill.background()
        label_box.line.fill.background()

        track_rect = row.get("trackRect") or {}
        fill_style = row.get("fillStyle") or {}
        track_style = row.get("trackStyle") or {}
        track_left = (block.layout.left or 0) + float(track_rect.get("x") or 0)
        track_top = (block.layout.top or 0) + float(track_rect.get("y") or 0)
        track_width = float(track_rect.get("w") or 0)
        track_height = float(track_rect.get("h") or 0)
        if track_width <= 0 or track_height <= 0:
            continue

        base_track = Block(
            kind="shape",
            layout=LayoutBox(left=track_left, top=track_top, width=track_width, height=track_height),
            shape_style={
                "fill_color": track_style.get("backgroundColor"),
                "border_radius": track_style.get("borderRadius"),
            },
        )
        add_shape_block(slide, base_track, prs, slide_model)

        ratio = float(row.get("ratio") or 0.0)
        fill_width = max(2.0, track_width * max(0.0, min(1.0, ratio)))
        fill_block = Block(
            kind="shape",
            layout=LayoutBox(left=track_left, top=track_top, width=fill_width, height=track_height),
            shape_style={
                "fill_color": fill_style.get("backgroundColor"),
                "border_radius": fill_style.get("borderRadius") or track_style.get("borderRadius"),
            },
        )
        add_shape_block(slide, fill_block, prs, slide_model)

        value_styles = row.get("valueStyles") or {}
        value_rect = row.get("valueRect") or {}
        value_box = slide.shapes.add_textbox(
            position_to_emu((block.layout.left or 0) + float(value_rect.get("x") or 0), "x", prs, slide_model),
            position_to_emu((block.layout.top or 0) + float(value_rect.get("y") or 0), "y", prs, slide_model),
            length_to_emu(float(value_rect.get("w") or chart.get("valueColumnWidth") or 80.0), "x", prs, slide_model),
            length_to_emu(float(value_rect.get("h") or track_height), "y", prs, slide_model),
        )
        set_text_frame_padding(value_box.text_frame)
        populate_text_frame(
            value_box.text_frame,
            [
                TextRun(
                    text=row.get("valueText", ""),
                    font_size=float(value_styles.get("fontSizePx") or label_style.get("font_size") or DEFAULT_FONT_SIZE * 0.65),
                    bold=css_weight_is_bold(value_styles.get("fontWeight")),
                    italic=(value_styles.get("fontStyle") or "").lower() == "italic",
                    color=value_styles.get("color"),
                    font_family=value_styles.get("fontFamily"),
                )
            ],
            {
                "font_size": float(value_styles.get("fontSizePx") or label_style.get("font_size") or DEFAULT_FONT_SIZE * 0.65),
                "bold": css_weight_is_bold(value_styles.get("fontWeight")),
                "italic": (value_styles.get("fontStyle") or "").lower() == "italic",
                "color": value_styles.get("color"),
                "align": (value_styles.get("textAlign") or "left").lower(),
                "font_family": value_styles.get("fontFamily"),
            },
            default_text=row.get("valueText", ""),
        )
        value_box.fill.background()
        value_box.line.fill.background()


def add_polyline_block(slide, block: Block, prs: Presentation, slide_model: SlideModel) -> None:
    points = block.vector_data.get("points") or []
    stroke_color = block.vector_data.get("stroke")
    stroke_width = float(block.vector_data.get("stroke_width") or 2.0)
    fill_color = block.vector_data.get("fill")
    closed = bool(block.vector_data.get("closed"))
    if len(points) < 2:
        return
    if closed and fill_color and fill_color.lower() != "none":
        # approximate closed polygon via Freeform
        builder = slide.shapes.build_freeform(
            position_to_emu(points[0]["x"], "x", prs, slide_model),
            position_to_emu(points[0]["y"], "y", prs, slide_model),
        )
        segments = [
            (
                position_to_emu(pt["x"], "x", prs, slide_model),
                position_to_emu(pt["y"], "y", prs, slide_model),
            )
            for pt in points[1:]
        ]
        builder.add_line_segments(segments, close=True)
        shape = builder.convert_to_shape()
        rgb_fill = css_color_to_rgb_tuple(fill_color)
        if rgb_fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(*rgb_fill)
        else:
            shape.fill.background()
        if stroke_color and stroke_color.lower() != "none":
            rgb = css_color_to_rgb_tuple(stroke_color)
            if rgb:
                shape.line.color.rgb = RGBColor(*rgb)
                shape.line.width = Pt(px_to_pt(stroke_width))
    else:
        for i in range(len(points) - 1):
            p1 = points[i]
            p2 = points[i + 1]
            connector = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                position_to_emu(p1["x"], "x", prs, slide_model),
                position_to_emu(p1["y"], "y", prs, slide_model),
                position_to_emu(p2["x"], "x", prs, slide_model),
                position_to_emu(p2["y"], "y", prs, slide_model),
            )
            if stroke_color and stroke_color.lower() != "none":
                rgb = css_color_to_rgb_tuple(stroke_color)
                if rgb:
                    connector.line.color.rgb = RGBColor(*rgb)
            connector.line.width = Pt(px_to_pt(stroke_width))


def add_circle_block(slide, block: Block, prs: Presentation, slide_model: SlideModel) -> None:
    circle = block.vector_data or {}
    cx = float(circle.get("cx") or 0.0)
    cy = float(circle.get("cy") or 0.0)
    r = float(circle.get("r") or 0.0)
    if r <= 0:
        return
    left = position_to_emu(cx - r, "x", prs, slide_model)
    top = position_to_emu(cy - r, "y", prs, slide_model)
    size = length_to_emu(r * 2, "x", prs, slide_model)
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, left, top, size, size)
    fill = circle.get("fill")
    stroke = circle.get("stroke")
    stroke_width = float(circle.get("strokeWidth") or 1.0)
    if fill and fill.lower() != "none":
        rgb = css_color_to_rgb_tuple(fill)
        if rgb:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(*rgb)
    else:
        shape.fill.background()
    if stroke and stroke.lower() != "none":
        rgb = css_color_to_rgb_tuple(stroke)
        if rgb:
            shape.line.color.rgb = RGBColor(*rgb)
            shape.line.width = Pt(px_to_pt(stroke_width))


def add_ellipse_block(slide, block: Block, prs: Presentation, slide_model: SlideModel) -> None:
    ellipse = block.vector_data or {}
    cx = float(ellipse.get("cx") or 0.0)
    cy = float(ellipse.get("cy") or 0.0)
    rx = float(ellipse.get("rx") or 0.0)
    ry = float(ellipse.get("ry") or 0.0)
    if rx <= 0 or ry <= 0:
        return
    left = position_to_emu(cx - rx, "x", prs, slide_model)
    top = position_to_emu(cy - ry, "y", prs, slide_model)
    width = length_to_emu(rx * 2, "x", prs, slide_model)
    height = length_to_emu(ry * 2, "y", prs, slide_model)
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, left, top, width, height)
    fill = ellipse.get("fill")
    stroke = ellipse.get("stroke")
    stroke_width = float(ellipse.get("strokeWidth") or 1.0)
    if fill and fill.lower() != "none":
        rgb = css_color_to_rgb_tuple(fill)
        if rgb:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(*rgb)
    else:
        shape.fill.background()
    if stroke and stroke.lower() != "none":
        rgb = css_color_to_rgb_tuple(stroke)
        if rgb:
            shape.line.color.rgb = RGBColor(*rgb)
            shape.line.width = Pt(px_to_pt(stroke_width))


def add_conic_gradient_block(slide, block: Block, prs: Presentation, slide_model: SlideModel) -> None:
    gradient = block.vector_data or {}
    segments = gradient.get("segments") or []
    if not segments:
        return
    cx = gradient.get("cx")
    cy = gradient.get("cy")
    radius = gradient.get("radius")
    if not radius:
        width = block.layout.width or 240
        height = block.layout.height or width
        radius = min(width, height) / 2
    if cx is None or cy is None:
        left = block.layout.left or DEFAULT_PADDING_X
        top = block.layout.top or DEFAULT_PADDING_Y
        width = block.layout.width or radius * 2
        height = block.layout.height or radius * 2
        cx = left + width / 2
        cy = top + height / 2
    for seg in segments:
        color = seg.get("color")
        start_deg = float(seg.get("startDeg") or 0.0)
        end_deg = float(seg.get("endDeg") or 0.0)
        if end_deg <= start_deg or not color:
            continue
        span = end_deg - start_deg
        steps = max(3, int(span / 8))
        def css_deg_to_point(deg):
            rad = math.radians(90 - deg)
            return (
                cx + radius * math.cos(rad),
                cy + radius * math.sin(rad),
            )
        builder = slide.shapes.build_freeform(
            position_to_emu(cx, "x", prs, slide_model),
            position_to_emu(cy, "y", prs, slide_model),
        )
        start_point = css_deg_to_point(start_deg)
        segments = [
            (
                position_to_emu(start_point[0], "x", prs, slide_model),
                position_to_emu(start_point[1], "y", prs, slide_model),
            )
        ]
        for i in range(1, steps + 1):
            angle = start_deg + (span * i / steps)
            point = css_deg_to_point(angle)
            segments.append(
                (
                    position_to_emu(point[0], "x", prs, slide_model),
                    position_to_emu(point[1], "y", prs, slide_model),
                )
            )
        builder.add_line_segments(segments, close=True)
        shape = builder.convert_to_shape()
        rgb = css_color_to_rgb_tuple(color)
        if rgb:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(*rgb)
        else:
            shape.fill.background()
        shape.line.fill.background()


def block_render_priority(block: Block) -> int:
    if block.kind in {"shape", "conic-gradient", "vector_circle", "vector_ellipse", "vector_polyline"}:
        return 0
    if block.kind in {"image", "table"}:
        return 1
    if block.kind == "bar-chart":
        return 2
    if block.kind in {"text", "list"}:
        return 3
    return 4


def slide_model_to_pptx(slides: List[SlideModel], output_path: str, page_size: str = "16:9") -> None:
    prs = Presentation()
    page_width, page_height = PPT_PAGE_SIZES_INCHES.get(page_size, PPT_PAGE_SIZES_INCHES["16:9"])
    prs.slide_width = Inches(page_width)
    prs.slide_height = Inches(page_height)
    blank = prs.slide_layouts[6]
    transform_mode = "fill" if page_size == "a4" else "contain"
    for slide_model in slides:
        slide_model.scale = None
        slide_model.scale_x = None
        slide_model.scale_y = None
        slide_model.offset_x = 0.0
        slide_model.offset_y = 0.0
        slide_model.transform_mode = transform_mode
        slide = prs.slides.add_slide(blank)
        if slide_model.background_color:
            rgb = css_color_to_rgb_tuple(slide_model.background_color)
            if rgb:
                fill = slide.background.fill
                fill.solid()
                fill.fore_color.rgb = RGBColor(*rgb)
        ensure_slide_transform(slide_model, prs)
        if add_rasterized_slide_block(slide, slide_model, prs):
            continue
        ordered_blocks = sorted(slide_model.blocks, key=lambda b: (b.z_index, block_render_priority(b), b.order))
        for block in ordered_blocks:
            try:
                if block.kind == "text":
                    add_text_block(slide, block, prs, slide_model)
                elif block.kind == "list":
                    add_list_block(slide, block, prs, slide_model)
                elif block.kind == "table":
                    add_table_block(slide, block, prs, slide_model)
                elif block.kind == "image":
                    add_image_block(slide, block, prs, slide_model)
                elif block.kind == "shape":
                    add_shape_block(slide, block, prs, slide_model)
                elif block.kind == "bar-chart":
                    add_bar_chart_block(slide, block, prs, slide_model)
                elif block.kind == "vector_polyline":
                    add_polyline_block(slide, block, prs, slide_model)
                elif block.kind == "vector_circle":
                    add_circle_block(slide, block, prs, slide_model)
                elif block.kind == "vector_ellipse":
                    add_ellipse_block(slide, block, prs, slide_model)
                elif block.kind == "conic-gradient":
                    add_conic_gradient_block(slide, block, prs, slide_model)
            except Exception as exc:
                print(f"[warn] スライド要素の描画に失敗しました: {exc}", file=sys.stderr)
    prs.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="HTML を編集可能な PPTX に変換します。")
    parser.add_argument("input_html", help="入力 HTML ファイルパス")
    parser.add_argument("output_pptx", help="出力 PPTX パス")
    parser.add_argument("--mode", default="editable", help="現在は 'editable' のみサポート")
    parser.add_argument(
        "--selector",
        default=".slide, .slide-container, [data-slide]",
        help="スライド要素として扱う CSS セレクタ（該当しない場合は <body> を1枚として扱う）",
    )
    parser.add_argument(
        "--engine",
        choices=["browser", "static"],
        default="browser",
        help="browser=Playwrightで実レイアウトを取得 / static=BeautifulSoupベースの簡易解析",
    )
    parser.add_argument(
        "--viewport",
        default="1920x1080",
        help="browser エンジン使用時のビューポート (例: 1366x768)",
    )
    parser.add_argument(
        "--dpr",
        type=int,
        default=2,
        help="browser エンジン使用時の device pixel ratio",
    )
    parser.add_argument(
        "--page-size",
        type=normalize_page_size,
        default="16:9",
        metavar="16:9|A4",
        help="出力 PPTX のページサイズ。16:9 または A4（横向き）。",
    )
    parser.add_argument(
        "--image-map",
        type=parse_image_map_entry,
        action="append",
        default=[],
        metavar="PLACEHOLDER=PATH",
        help="画像 placeholder を実ファイルにマッピング。例: --image-map IMAGE_URL_1=./assets/hero.png",
    )
    parser.add_argument(
        "--rasterize-slides",
        default="none",
        metavar="auto|none|3-6,8",
        help="browser エンジン時に指定 slide をページ画像として埋め込む。既定は none。auto / none / 1始まりの番号・範囲指定。",
    )
    args = parser.parse_args()
    image_map = build_image_map(args.image_map)

    input_path = Path(args.input_html)
    if not input_path.exists():
        print(f"入力ファイルが存在しません: {input_path}", file=sys.stderr)
        sys.exit(1)
    if args.mode != "editable":
        print("サポートされていないモードです。--mode editable を利用してください。", file=sys.stderr)
        sys.exit(2)

    try:
        if "x" in args.viewport.lower():
            vw_str, vh_str = args.viewport.lower().split("x", 1)
            viewport_w = int(vw_str)
            viewport_h = int(vh_str)
        else:
            viewport_w = 1920
            viewport_h = 1080
        if args.engine == "browser":
            slides = parse_html_browser(
                str(input_path),
                selector=args.selector,
                viewport_w=viewport_w,
                viewport_h=viewport_h,
                dpi_scale=args.dpr,
                image_map=image_map,
                rasterize_slides=args.rasterize_slides,
            )
        else:
            slides = parse_html_static(str(input_path), image_map=image_map)
        slide_model_to_pptx(slides, args.output_pptx, page_size=args.page_size)
    except Exception as exc:
        print(f"変換に失敗しました: {exc}", file=sys.stderr)
        sys.exit(3)


if __name__ == "__main__":
    main()
